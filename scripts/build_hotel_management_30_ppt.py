from __future__ import annotations

from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
BASE = ROOT / "Aurelia_Collection_Hotel_Management_Project_Full_Deck.pptx"
OUT = ROOT / "Aurelia_Collection_Hotel_Management_Project_30_Slides.pptx"
NAVY = RGBColor(11, 20, 39); NAVY2 = RGBColor(25, 52, 92); NAVY3 = RGBColor(43, 77, 119)
GOLD = RGBColor(195, 161, 90); PALE_GOLD = RGBColor(243, 227, 184); IVORY = RGBColor(250, 248, 244)
BG = RGBColor(247, 248, 250); WHITE = RGBColor(255, 255, 255); INK = RGBColor(15, 23, 42)
MUTED = RGBColor(100, 116, 139); GREEN = RGBColor(6, 118, 71); BLUE = RGBColor(0, 122, 255)
LINE = RGBColor(226, 232, 240); LIGHT_GREEN = RGBColor(236, 253, 243); LIGHT_BLUE = RGBColor(230, 240, 252)
prs = Presentation(str(BASE)); blank = prs.slide_layouts[6]

def shape(slide, kind, x, y, w, h, fill=None, line=None, transparency=0):
    shp=slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb=fill; shp.fill.transparency=transparency
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb=line; shp.line.width=Pt(.8)
    return shp

def rounded(slide,x,y,w,h,fill=WHITE,line=None): return shape(slide,MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,x,y,w,h,fill,line)
def rect(slide,x,y,w,h,fill,line=None): return shape(slide,MSO_AUTO_SHAPE_TYPE.RECTANGLE,x,y,w,h,fill,line)
def circle(slide,x,y,d,fill,line=None): return shape(slide,MSO_AUTO_SHAPE_TYPE.OVAL,x,y,d,d,fill,line)

def text(slide,value,x,y,w,h,size=14,color=INK,bold=False,font="Aptos",align=PP_ALIGN.LEFT,valign=MSO_ANCHOR.TOP,margin=.04):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=Inches(margin); tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=value; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return box

def bullets(slide,items,x,y,w,h,size=11,color=INK,gap=5,bullet_color=GOLD):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.04); tf.margin_top=tf.margin_bottom=Inches(.02)
    for i,item in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.space_after=Pt(gap); p.text=""
        r=p.add_run(); r.text="•  "; r.font.name="Aptos"; r.font.size=Pt(size); r.font.bold=True; r.font.color.rgb=bullet_color
        r2=p.add_run(); r2.text=item; r2.font.name="Aptos"; r2.font.size=Pt(size); r2.font.color.rgb=color
    return box

def header(slide,kicker,title,subtitle,num,dark=False):
    text(slide,kicker.upper(),.62,.43,5.3,.22,9.5,GOLD,True,margin=0)
    text(slide,title,.62,.70,11.7,.58,27,WHITE if dark else INK,True,"Georgia",margin=0)
    text(slide,subtitle,.64,1.32,11.3,.34,11.5,RGBColor(205,216,232) if dark else MUTED,margin=0)
    text(slide,f"{num:02d}",12.35,.45,.4,.22,9,RGBColor(205,216,232) if dark else MUTED,True,align=PP_ALIGN.RIGHT,margin=0)

def footer(slide,num,dark=False):
    c=RGBColor(185,198,218) if dark else RGBColor(148,163,184); rect(slide,.62,7.13,12.08,.008,GOLD if dark else LINE)
    text(slide,"Aurelia Collection · Hotel management project · 28 September 2026",.64,7.19,8.9,.18,8.2,c,margin=0)
    text(slide,f"{num:02d}",12.15,7.16,.55,.22,8.5,c,True,align=PP_ALIGN.RIGHT,margin=0)

def bg(slide,color=BG): slide.background.fill.solid(); slide.background.fill.fore_color.rgb=color

def card(slide,x,y,w,h,title,desc,fill=WHITE,accent=GOLD,title_color=INK,desc_color=MUTED):
    rounded(slide,x,y,w,h,fill,LINE if fill!=NAVY else None); rect(slide,x,y,.055,h,accent)
    text(slide,title,x+.22,y+.20,w-.35,.24,13,title_color,True,"Georgia",margin=0)
    text(slide,desc,x+.22,y+.62,w-.35,h-.70,10.2,desc_color,margin=0)

# 23 — users
s=prs.slides.add_slide(blank); bg(s,IVORY); header(s,"Users","Who uses Aurelia Collection?","The same platform serves guests, front office, room operations, management and system administration.",23)
users=[("Guest","Books and manages a personal stay","Public website, My Stays and self-service cancellation",GOLD),("Reception","Converts demand into stays","Walk-in, phone, OTA/agency, arrivals and departures",NAVY3),("Housekeeping","Keeps rooms ready to sell","Room Board, tasks, room status and maintenance",GREEN),("Manager","Controls commercial performance","Rates, finance reports, analytics and property scope",BLUE),("Administrator","Maintains the platform","Users, configuration, Django admin and all manager access",NAVY)]
for i,(a,b,c,col) in enumerate(users):
    x=.78+(i%3)*4.13; y=1.92+(i//3)*1.75
    card(s,x,y,3.63,1.28,a,b+"\n"+c,WHITE,col)
rounded(s,.78,5.70,11.84,.65,NAVY,None); text(s,"Simple rule: guests see their own journey; staff operate the hotel; managers control business; administrators control the system.",1.05,5.91,11.3,.18,11,WHITE,True,align=PP_ALIGN.CENTER,margin=0); footer(s,23)

# 24 — website service
s=prs.slides.add_slide(blank); bg(s,BG); header(s,"Guest service","Website service: from interest to confirmed stay","The website is the direct channel for discovery, booking, payment intent and post-booking self-service.",24)
steps=[("1","Discover","Properties, room photos, location, Google Maps and collection recommendations."),("2","Search","Dates, guests, rooms and live availability."),("3","Choose","Room type, amenities, nightly rates and full-stay price."),("4","Offer","Registered first-time website guest can use the one-time 10% offer."),("5","Confirm","Contact details, special requests, payment method and confirmation."),("6","Manage","My Stays, booking details and cancellation when allowed.")]
for i,(n,t,d) in enumerate(steps):
    x=.78+(i%3)*4.13; y=1.92+(i//3)*2.0; rounded(s,x,y,3.63,1.48,WHITE,LINE); circle(s,x+.24,y+.22,.40,NAVY if i in (0,5) else GOLD); text(s,n,x+.24,y+.33,.40,.12,9,WHITE if i in (0,5) else NAVY,True,align=PP_ALIGN.CENTER,margin=0); text(s,t,x+.84,y+.24,2.3,.20,13,NAVY,True,"Georgia",margin=0); text(s,d,x+.24,y+.80,3.05,.42,9.8,MUTED,margin=0)
rounded(s,.78,6.20,11.84,.50,LIGHT_GREEN,RGBColor(167,243,208)); text(s,"Website result: one validated booking, one folio, one payment record and one protected financial snapshot.",1.02,6.37,11.35,.16,10,GREEN,True,align=PP_ALIGN.CENTER,margin=0); footer(s,24)

# 25 — walk-in / phone
s=prs.slides.add_slide(blank); bg(s,IVORY); header(s,"Front office","Walk-in and phone service","Reception creates the reservation for the guest and keeps the same inventory, pricing and folio controls.",25)
card(s,.78,1.92,5.70,3.85,"Walk-in · source: walk_in","1. Identify the guest\n2. Select room type and dates\n3. Validate availability and price\n4. Create reservation without requiring an online account\n5. Record cash, card, wallet or bank transfer\n6. Check in when a vacant-clean room is assigned",WHITE,GOLD)
card(s,6.84,1.92,5.78,3.85,"Phone · source: phone","1. Take the caller’s dates and guest details\n2. Search room availability\n3. Quote the full-stay total\n4. Create the booking for the caller\n5. Keep contact details and requests on the record\n6. Send or explain the confirmation",WHITE,NAVY3)
rounded(s,.78,6.08,11.84,.62,NAVY,None); text(s,"Front desk principle: every manual booking uses the same validation and cannot bypass room inventory or financial snapshots.",1.04,6.29,11.34,.17,10.5,WHITE,True,align=PP_ALIGN.CENTER,margin=0); footer(s,25)

# 26 — OTA agency
s=prs.slides.add_slide(blank); bg(s,BG); header(s,"Distribution","OTA and agency service","External demand is recorded in the PMS today; automatic channel-manager synchronization is the planned next layer.",26)
rounded(s,.78,1.92,4.10,4.55,NAVY,None); text(s,"TODAY",1.12,2.30,1.1,.20,10,PALE_GOLD,True,margin=0); text(s,"Manual but controlled",1.12,2.72,3.1,.30,20,WHITE,True,"Georgia",margin=0); bullets(s,["Reception selects OTA / Agency.","Guest, dates, room and source are recorded.","External reference can be kept in notes.","Availability, pricing and folio stay inside Aurelia.","Channel source remains available for reporting."],1.12,3.45,3.25,1.9,11,RGBColor(215,226,239),8,PALE_GOLD)
rounded(s,5.28,1.92,7.34,4.55,WHITE,LINE); text(s,"NEXT",5.66,2.30,1.1,.20,10,GREEN,True,margin=0); text(s,"Automatic channel connection",5.66,2.72,4.6,.30,20,NAVY,True,"Georgia",margin=0); bullets(s,["Channel-manager API for availability and rates.","Booking push webhooks from OTA partners.","Automatic reservation import and source mapping.","Commission and channel-specific rate plans.","Overbooking protection across connected channels."],5.66,3.45,5.95,1.9,11,INK,8,GREEN)
rounded(s,5.66,5.74,6.35,.42,LIGHT_BLUE,RGBColor(191,219,254)); text(s,"Current boundary: no live Booking.com/Expedia-style synchronization yet.",5.90,5.87,5.88,.14,9.5,NAVY3,True,align=PP_ALIGN.CENTER,margin=0); footer(s,26)

# 27 — task matrix
s=prs.slides.add_slide(blank); bg(s,IVORY); header(s,"Role playbook","Which role handles each task?","A practical handoff view for training and daily accountability.",27)
cols=[.78,3.15,5.60,8.05,10.50]; widths=[2.18,2.18,2.18,2.18,2.12]
for x,w,l in zip(cols,widths,["TASK","RECEPTION","HOUSEKEEPING","MANAGER","ADMIN"]): rounded(s,x,1.92,w,.44,NAVY,None); text(s,l,x+.10,2.06,w-.20,.13,8.5,PALE_GOLD,True,margin=0)
rows=[("Create reservation","✓","—","Review","Override"),("Check in / out","✓","—","Review","Override"),("Update room status","✓","✓","Review","Override"),("Complete room task","—","✓","Review","Override"),("Set rate rules","—","—","✓","Override"),("View reports","—","—","✓","✓"),("Manage users","—","—","—","✓")]
for i,row in enumerate(rows):
    y=2.38+i*.54; fill=WHITE if i%2==0 else RGBColor(246,248,251)
    for x,w in zip(cols,widths): rect(s,x,y,w,.51,fill,LINE)
    text(s,row[0],cols[0]+.10,y+.16,widths[0]-.18,.15,9.8,INK,True,margin=0)
    for j in range(1,5):
        c=GREEN if row[j]=="✓" else (NAVY3 if row[j] in ("Review","Override") else MUTED)
        text(s,row[j],cols[j],y+.16,widths[j],.15,9.5,c,True,align=PP_ALIGN.CENTER,margin=0)
rounded(s,.78,6.30,11.84,.42,LIGHT_BLUE,RGBColor(191,219,254)); text(s,"Current technical gate: all staff roles share the broad staff gate; manager/admin controls rates, reporting and system administration.",1.02,6.43,11.35,.14,9.3,NAVY3,True,align=PP_ALIGN.CENTER,margin=0); footer(s,27)

# 28 — pricing example
s=prs.slides.add_slide(blank); bg(s,NAVY); header(s,"Pricing example","A Cambodia-friendly price that is easy to explain","The current demo starts at $30/night and uses 5% tax only — no separate service charge.",28,True)
rounded(s,.78,1.92,5.15,4.58,RGBColor(23,40,68),RGBColor(69,92,125)); text(s,"Example: entry room",1.12,2.28,2.6,.22,14,PALE_GOLD,True,"Georgia",margin=0); text(s,"$30",1.12,2.78,2.1,.56,36,WHITE,True,"Georgia",margin=0); text(s,"per night",3.18,3.02,1.25,.18,11,RGBColor(200,214,233),margin=0)
calc=[("2 nights","$60.00"),("Welcome offer 10%","−$6.00"),("Tax-only subtotal","$54.00"),("Tax 5%","+$2.70"),("Guest total","$56.70")]
for i,(a,b) in enumerate(calc):
    y=3.72+i*.43; text(s,a,1.12,y,2.35,.18,10.5,RGBColor(205,216,232) if i<4 else PALE_GOLD,i==4,margin=0); text(s,b,4.15,y,1.12,.18,10.5,WHITE if i<4 else PALE_GOLD,i==4,align=PP_ALIGN.RIGHT,margin=0)
rounded(s,6.30,1.92,6.32,4.58,PALE_GOLD,None); text(s,"Current catalog ladder",6.66,2.28,2.9,.22,14,NAVY,True,"Georgia",margin=0); text(s,"$30 → $460",6.66,2.72,3.2,.42,28,NAVY,True,"Georgia",margin=0); bullets(s,["Entry rooms: $30–$58","Suites and city rooms: $64–$172","Villas and premium stays: $216–$356","Highest current catalog rate: $460","Historical old bookings keep their original totals"],6.66,3.52,5.2,1.62,11,NAVY,8,GREEN)
text(s,"New prices affect the current catalog; they do not rewrite historical financial snapshots.",6.66,5.74,5.2,.28,10,RGBColor(85,67,27),True,margin=0); footer(s,28,True)

# 29 — controls
s=prs.slides.add_slide(blank); bg(s,BG); header(s,"Controls","The safeguards behind the simple experience","The interface is easy because the business rules are controlled underneath.",29)
controls=[("Availability","Checks every night and prevents overselling.",NAVY),("Pricing","Stores rate, discount, tax and total on the booking.",GOLD),("Rooms","Only vacant-clean rooms can be assigned at check-in.",GREEN),("Maintenance","High/critical rooms are blocked from sale.",BLUE),("Cancellation","Applies policy and creates refund entries when needed.",NAVY3),("Audit","Tests money, inventory, folios and room states.",RGBColor(130,74,20))]
for i,(t,d,c) in enumerate(controls):
    x=.78+(i%3)*4.13; y=1.92+(i//3)*1.74; card(s,x,y,3.63,1.30,t,d,WHITE,c)
rounded(s,.78,5.72,11.84,.64,NAVY,None); text(s,"Verified release: 23/23 tests · 2,160 ledger invariants · 0 pending migrations · public routes checked",1.02,5.94,11.35,.18,10.5,WHITE,True,align=PP_ALIGN.CENTER,margin=0); footer(s,29)

# 30 — demo and handoff
s=prs.slides.add_slide(blank); bg(s,IVORY); header(s,"Demonstration","The clearest way to present the project","Use this sequence to explain the system to an owner, hotel team or technical reviewer.",30)
seq=[("1","Open public home","Show the signed-out guest experience."),("2","Choose a property","Show curated collection and room discovery."),("3","Search $30 room","Show dates, availability and 5% tax-only pricing."),("4","Book as website guest","Show welcome offer and confirmation."),("5","Open PMS","Show dashboard, arrivals and reservations."),("6","Create walk-in","Show source selection and front-desk control."),("7","Operate the room","Show check-in, Room Board and housekeeping."),("8","Close the stay","Show checkout, folio, audit and manager report.")]
for i,(n,t,d) in enumerate(seq):
    x=.78+(i%4)*3.08; y=1.92+(i//4)*1.73; rounded(s,x,y,2.67,1.30,WHITE,LINE); circle(s,x+.20,y+.20,.34,GOLD if i<4 else NAVY3); text(s,n,x+.20,y+.29,.34,.11,8.5,NAVY if i<4 else WHITE,True,align=PP_ALIGN.CENTER,margin=0); text(s,t,x+.68,y+.20,1.75,.18,10.6,INK,True,margin=0); text(s,d,x+.20,y+.70,2.18,.30,9.2,MUTED,margin=0)
rounded(s,.78,5.90,11.84,.70,NAVY,None); text(s,"Next handoff: configure and verify the owner’s MySQL/Navicat connection, then add live OTA synchronization when ready.",1.04,6.14,11.34,.20,11,WHITE,True,align=PP_ALIGN.CENTER,margin=0); footer(s,30)

prs.core_properties.title="Aurelia Collection — Hotel Management Project · 30-slide latest version"
prs.core_properties.subject="Clear end-to-end hotel project explanation, roles, channels, pricing, operations and handoff"
prs.core_properties.author="Aurelia Collection"
prs.core_properties.keywords="hotel management, PMS, Khmer, roles, website, walk-in, OTA, pricing, operations"
prs.save(OUT)
print(OUT)
