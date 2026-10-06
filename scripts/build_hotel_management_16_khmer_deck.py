from __future__ import annotations

from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
BASE = ROOT / "Aurelia_Collection_Hotel_Management_Project_Short_Deck_KH.pptx"
OUT = ROOT / "Aurelia_Collection_Hotel_Management_Project_16_Slides_KH.pptx"
prs = Presentation(str(BASE))
blank = prs.slide_layouts[6]

NAVY = RGBColor(11, 20, 39)
NAVY2 = RGBColor(25, 52, 92)
NAVY3 = RGBColor(43, 77, 119)
GOLD = RGBColor(195, 161, 90)
PALE_GOLD = RGBColor(243, 227, 184)
IVORY = RGBColor(250, 248, 244)
BG = RGBColor(247, 248, 250)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(15, 23, 42)
MUTED = RGBColor(100, 116, 139)
GREEN = RGBColor(6, 118, 71)
BLUE = RGBColor(37, 99, 235)
RED = RGBColor(180, 35, 24)
LINE = RGBColor(226, 232, 240)
LIGHT_GREEN = RGBColor(236, 253, 243)
LIGHT_BLUE = RGBColor(239, 246, 255)
LIGHT_GOLD = RGBColor(246, 239, 223)


def shape(slide, kind, x, y, w, h, fill=None, line=None):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(0.8)
    return shp


def rounded(slide, x, y, w, h, fill=WHITE, line=None):
    return shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line)


def rect(slide, x, y, w, h, fill, line=None):
    return shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, w, h, fill, line)


def circle(slide, x, y, d, fill, line=None):
    return shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, y, d, d, fill, line)


def text(slide, value, x, y, w, h, size=14, color=INK, bold=False,
         font="Noto Sans Khmer", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = value; r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=10.5, color=INK, gap=5, bullet_color=GOLD):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(gap); p.text = ""
        r = p.add_run(); r.text = "•  "; r.font.name = "Noto Sans Khmer"; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = bullet_color
        r2 = p.add_run(); r2.text = item; r2.font.name = "Noto Sans Khmer"; r2.font.size = Pt(size); r2.font.color.rgb = color
    return box


def bg(slide, color=BG):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = color


def header(slide, kicker, title, subtitle, num, dark=False):
    text(slide, kicker, .62, .38, 5.5, .25, 9.5, GOLD, True, margin=0)
    text(slide, title, .62, .69, 11.7, .62, 25, WHITE if dark else INK, True, "Noto Sans Khmer", margin=0)
    text(slide, subtitle, .64, 1.33, 11.4, .38, 10.7, RGBColor(205,216,232) if dark else MUTED, margin=0)
    text(slide, f"{num:02d}", 12.35, .40, .40, .22, 9, RGBColor(205,216,232) if dark else MUTED, True, align=PP_ALIGN.RIGHT, margin=0)


def footer(slide, num, dark=False):
    c = RGBColor(70,92,125) if dark else LINE
    tc = RGBColor(185,198,218) if dark else RGBColor(148,163,184)
    rect(slide, .62, 7.10, 12.08, .008, c)
    text(slide, "Aurelia Collection · គម្រោងគ្រប់គ្រងសណ្ឋាគារ · ២៩ កញ្ញា ២០២៦", .64, 7.16, 10.0, .18, 8.0, tc, margin=0)
    text(slide, f"{num:02d}", 12.15, 7.13, .55, .22, 8.5, tc, True, align=PP_ALIGN.RIGHT, margin=0)


def card(slide, x, y, w, h, title, desc, fill=WHITE, accent=GOLD, title_size=12.5, desc_size=9.8):
    rounded(slide, x, y, w, h, fill, LINE if fill != NAVY else None)
    rect(slide, x, y, .055, h, accent)
    text(slide, title, x+.22, y+.18, w-.35, .28, title_size, INK if fill != NAVY else WHITE, True, margin=0)
    text(slide, desc, x+.22, y+.60, w-.35, h-.68, desc_size, MUTED if fill != NAVY else RGBColor(205,216,232), margin=0)


def flow(slide, x, y, n, title, desc, accent=GOLD):
    rounded(slide, x, y, 2.72, 1.45, WHITE, LINE)
    circle(slide, x+.20, y+.20, .38, accent)
    text(slide, str(n), x+.20, y+.30, .38, .10, 8.5, NAVY if accent == GOLD else WHITE, True, align=PP_ALIGN.CENTER, margin=0)
    text(slide, title, x+.72, y+.19, 1.78, .25, 10.8, INK, True, margin=0)
    text(slide, desc, x+.20, y+.76, 2.28, .45, 8.9, MUTED, margin=0)

# 11 — guest website detail
s = prs.slides.add_slide(blank); bg(s, BG)
header(s, "លំហូរភ្ញៀវ", "ការកក់តាមគេហទំព័រ ក្នុង ៦ ជំហាន", "ភ្ញៀវឃើញផ្លូវដែលច្បាស់ ខណៈប្រព័ន្ធធ្វើការផ្ទៀងផ្ទាត់ស្មុគស្មាញនៅខាងក្រោយ។", 11)
steps = [
    ("ស្វែងរក", "មើលសណ្ឋាគារ, ទីតាំង, រូបថត និងបរិយាកាស។"),
    ("ជ្រើសកាលបរិច្ឆេទ", "ជ្រើសថ្ងៃស្នាក់នៅ, ចំនួនភ្ញៀវ និងបន្ទប់។"),
    ("ប្រៀបធៀប", "មើលប្រភេទបន្ទប់, សេវា និងតម្លៃពេញការស្នាក់នៅ។"),
    ("Welcome Offer", "អ្នកចុះឈ្មោះដែលមានសិទ្ធិ ទទួលបាន ១០% ម្តង។"),
    ("បញ្ជាក់", "បញ្ចូលព័ត៌មានទំនាក់ទំនង, សំណើ និងវិធីទូទាត់។"),
    ("My Stays", "មើលការកក់ និងគ្រប់គ្រងការស្នាក់នៅដោយខ្លួនឯង។"),
]
for i, (t, d) in enumerate(steps):
    flow(s, .78 + (i % 3) * 4.13, 1.95 + (i // 3) * 1.92, i+1, t, d, GOLD if i in (0,5) else NAVY3)
rounded(s, .78, 6.12, 11.84, .54, LIGHT_GREEN, RGBColor(167,243,208))
text(s, "លទ្ធផល: ការកក់តែមួយ ដែលមាន Availability, តម្លៃ, Folio និងប្រវត្តិច្បាស់លាស់។", 1.03, 6.31, 11.34, .18, 10, GREEN, True, align=PP_ALIGN.CENTER, margin=0)
footer(s, 11)

# 12 — walk-in / phone
s = prs.slides.add_slide(blank); bg(s, IVORY)
header(s, "Front Office", "សេវា Walk-in និងទូរស័ព្ទ", "Reception បង្កើតការកក់ជំនួសភ្ញៀវ ដោយប្រើ Availability, Pricing និង Folio ដូចគ្នានឹង Website។", 12)
card(s, .78, 1.95, 5.70, 3.95, "Walk-in · source: walk_in", "១. ស្វាគមន៍ និងកំណត់តម្រូវការភ្ញៀវ\n២. ពិនិត្យបន្ទប់ និងកាលបរិច្ឆេទ\n៣. បង្ហាញតម្លៃសរុបឱ្យច្បាស់\n៤. បង្កើតការកក់ និងកត់ត្រាប្រភព\n៥. កត់ត្រា Cash, Card ឬ Transfer\n៦. Check-in នៅពេលបន្ទប់រួចរាល់", WHITE, GOLD, 12.5, 9.8)
card(s, 6.84, 1.95, 5.78, 3.95, "Phone · source: phone", "១. ស្តាប់ថ្ងៃស្នាក់នៅ និងព័ត៌មានភ្ញៀវ\n២. ស្វែងរក Availability\n៣. ពន្យល់តម្លៃពេញការស្នាក់នៅ\n៤. បង្កើតការកក់សម្រាប់អ្នកហៅ\n៥. រក្សា Contact និងសំណើពិសេស\n៦. ផ្ញើ ឬអានលេខ Confirmation", WHITE, NAVY3, 12.5, 9.8)
rounded(s, .78, 6.18, 11.84, .54, NAVY, None)
text(s, "គោលការណ៍ Front Desk: ការកក់ដោយដៃ ក៏ត្រូវឆ្លងកាត់ច្បាប់តម្លៃ និងការពារស្តុកដូចគ្នា។", 1.02, 6.37, 11.35, .17, 9.8, WHITE, True, align=PP_ALIGN.CENTER, margin=0)
footer(s, 12, True)

# 13 — OTA / agency
s = prs.slides.add_slide(blank); bg(s, BG)
header(s, "Distribution", "សេវា OTA និង Agency", "តម្រូវការពីខាងក្រៅ ត្រូវបានកត់ត្រាក្នុង PMS បច្ចុប្បន្ន; Channel Manager គឺជាជំហានបន្ទាប់។", 13)
rounded(s, .78, 1.95, 4.10, 4.65, NAVY, None)
text(s, "បច្ចុប្បន្ន", 1.12, 2.30, 1.2, .22, 10, PALE_GOLD, True, margin=0)
text(s, "Manual ប៉ុន្តែមានការគ្រប់គ្រង", 1.12, 2.72, 3.2, .34, 18, WHITE, True, margin=0)
bullets(s, ["Reception ជ្រើស OTA / Agency។", "កត់ត្រាភ្ញៀវ, ថ្ងៃ, បន្ទប់ និងប្រភព។", "រក្សាលេខយោងខាងក្រៅក្នុង Notes។", "Availability, Pricing និង Folio នៅក្នុង Aurelia។", "ប្រភព Channel មានសម្រាប់របាយការណ៍។"], 1.12, 3.48, 3.25, 2.0, 10.4, RGBColor(215,226,239), 8, PALE_GOLD)
rounded(s, 5.28, 1.95, 7.34, 4.65, WHITE, LINE)
text(s, "ជំហានបន្ទាប់", 5.66, 2.30, 1.6, .22, 10, GREEN, True, margin=0)
text(s, "Automatic Channel Connection", 5.66, 2.72, 4.7, .34, 18, NAVY, True, margin=0)
bullets(s, ["API សម្រាប់ Availability និង Rate។", "Import ការកក់ពីដៃគូ OTA។", "ផ្គូផ្គងប្រភពដោយស្វ័យប្រវត្តិ។", "Rate Plan និង Commission តាម Channel។", "ការពារ Overbooking រវាង Channel។"], 5.66, 3.48, 5.95, 2.0, 10.4, INK, 8, GREEN)
rounded(s, 5.66, 5.92, 6.35, .40, LIGHT_BLUE, RGBColor(191,219,254))
text(s, "មិនទាន់មាន Live Synchronization ជាមួយ Booking.com / Expedia នៅឡើយទេ។", 5.90, 6.04, 5.88, .15, 9.0, NAVY3, True, align=PP_ALIGN.CENTER, margin=0)
footer(s, 13)

# 14 — responsibility matrix
s = prs.slides.add_slide(blank); bg(s, IVORY)
header(s, "តួនាទី", "អ្នកណាទទួលខុសត្រូវលើកិច្ចការណា?", "ផែនទីខ្លីសម្រាប់បណ្តុះបណ្តាល និងការប្រគល់ការងារប្រចាំថ្ងៃ។", 14)
cols = [.78, 3.15, 5.60, 8.05, 10.50]; widths = [2.18, 2.18, 2.18, 2.18, 2.12]
for x, w, label in zip(cols, widths, ["កិច្ចការ", "RECEPTION", "HOUSEKEEPING", "MANAGER", "ADMIN"]):
    rounded(s, x, 1.95, w, .44, NAVY, None)
    text(s, label, x+.10, 2.07, w-.20, .14, 8.2, PALE_GOLD, True, align=PP_ALIGN.CENTER if x != .78 else PP_ALIGN.LEFT, margin=0)
rows = [
    ("បង្កើតការកក់", "✓", "—", "ពិនិត្យ", "ពេញសិទ្ធិ"),
    ("Check-in / Check-out", "✓", "—", "ពិនិត្យ", "ពេញសិទ្ធិ"),
    ("កែស្ថានភាពបន្ទប់", "✓", "✓", "ពិនិត្យ", "ពេញសិទ្ធិ"),
    ("បញ្ចប់កិច្ចការ", "—", "✓", "ពិនិត្យ", "ពេញសិទ្ធិ"),
    ("កំណត់ Rate Rule", "—", "—", "✓", "ពេញសិទ្ធិ"),
    ("មើលរបាយការណ៍", "—", "—", "✓", "✓"),
    ("គ្រប់គ្រង Users", "—", "—", "—", "✓"),
]
for i, row in enumerate(rows):
    y = 2.39 + i * .54; fill = WHITE if i % 2 == 0 else RGBColor(246,248,251)
    for x, w in zip(cols, widths): rect(s, x, y, w, .51, fill, LINE)
    text(s, row[0], cols[0]+.10, y+.16, widths[0]-.18, .15, 9.0, INK, True, margin=0)
    for j in range(1, 5):
        c = GREEN if row[j] == "✓" else (NAVY3 if row[j] in ("ពិនិត្យ", "ពេញសិទ្ធិ") else MUTED)
        text(s, row[j], cols[j], y+.16, widths[j], .15, 8.8, c, True, align=PP_ALIGN.CENTER, margin=0)
rounded(s, .78, 6.30, 11.84, .42, LIGHT_BLUE, RGBColor(191,219,254))
text(s, "គោលការណ៍: អ្នកប្រើឃើញតែការងារដែលខ្លួនទទួលខុសត្រូវ ហើយប្រព័ន្ធការពារស្តុក និងហិរញ្ញវត្ថុរួម។", 1.02, 6.43, 11.35, .14, 9.2, NAVY3, True, align=PP_ALIGN.CENTER, margin=0)
footer(s, 14)

# 15 — financial integrity
s = prs.slides.add_slide(blank); bg(s, NAVY)
header(s, "ហិរញ្ញវត្ថុ", "លេខគ្រប់ចំនួន ត្រូវមានប្រភពច្បាស់", "ការកក់មួយៗរក្សាតម្លៃ និងចំនួនសរុបដែលបានលក់ពិតប្រាកដ ដើម្បីការពារប្រវត្តិហិរញ្ញវត្ថុ។", 15, True)
rounded(s, .78, 1.95, 5.25, 4.62, RGBColor(23,40,68), RGBColor(69,92,125))
text(s, "គោលការណ៍ Snapshot", 1.12, 2.30, 2.8, .25, 14, PALE_GOLD, True, margin=0)
bullets(s, ["Booking រក្សា Rate Detail និងចំនួនសរុប។", "Discount, Tax និង Total ត្រូវបានរក្សាជាមួយការកក់។", "ការកែ Catalog ថ្មី មិនកែចំណូលប្រវត្តិសាស្ត្រ។", "Folio និង Payment អាចផ្ទៀងផ្ទាត់ត្រឡប់ទៅ Booking។", "Cancellation និង Refund មានកំណត់ត្រាច្បាស់។"], 1.12, 2.88, 4.22, 2.18, 10.6, RGBColor(215,226,239), 9, PALE_GOLD)
rounded(s, 6.38, 1.95, 6.24, 4.62, PALE_GOLD, None)
text(s, "ឧទាហរណ៍ការគណនា", 6.76, 2.30, 2.8, .25, 14, NAVY, True, margin=0)
calc = [("២ យប់ × $30", "$60.00"), ("បញ្ចុះ Welcome 10%", "−$6.00"), ("Subtotal", "$54.00"), ("ពន្ធ 5%", "+$2.70"), ("សរុបភ្ញៀវ", "$56.70")]
for i, (a, b) in enumerate(calc):
    y = 2.95 + i * .48
    text(s, a, 6.76, y, 3.20, .20, 10.5, NAVY, i == 4, margin=0)
    text(s, b, 11.05, y, 1.12, .20, 10.5, NAVY, i == 4, align=PP_ALIGN.RIGHT, margin=0)
text(s, "Service Charge: 0% · Tax: 5%", 6.76, 5.62, 4.8, .20, 10.4, RGBColor(85,67,27), True, margin=0)
footer(s, 15, True)

# 16 — handoff
s = prs.slides.add_slide(blank); bg(s, IVORY)
header(s, "ការប្រគល់គម្រោង", "អ្វីដែលត្រូវធ្វើបន្ទាប់", "កំណែបច្ចុប្បន្នត្រៀមសម្រាប់ Demo និងប្រើប្រាស់; ជំហានបន្ទាប់គឺការភ្ជាប់បរិស្ថានរបស់ម្ចាស់ប្រព័ន្ធ។", 16)
card(s, .78, 1.95, 3.63, 2.02, "១. Database", "កំណត់ និងផ្ទៀងផ្ទាត់ការភ្ជាប់ MySQL/MariaDB តាម Navicat ដោយប្រើ Credentials ឯកជនរបស់ម្ចាស់ប្រព័ន្ធ។", WHITE, GOLD, 12.5, 9.8)
card(s, 4.84, 1.95, 3.63, 2.02, "២. Distribution", "បន្ថែម Channel Manager ដើម្បីធ្វើ Availability, Rate និងការនាំចូល Booking ពី OTA ដោយស្វ័យប្រវត្តិ។", WHITE, BLUE, 12.5, 9.8)
card(s, 8.90, 1.95, 3.72, 2.02, "៣. Maturity", "បន្ថែម Folio Line Items សម្រាប់ Minibar, Spa, Restaurant និង Transport តាមតម្រូវការសណ្ឋាគារ។", WHITE, GREEN, 12.5, 9.8)
text(s, "លំដាប់ Demo ដែលណែនាំ", .82, 4.58, 3.2, .25, 17, NAVY, True, margin=0)
bullets(s, ["បើក Website និងស្វែងរកបន្ទប់ $30។", "បញ្ជាក់ការកក់ និងបង្ហាញ My Stays។", "បើក PMS ដើម្បីបង្ហាញ Reservation និង Dashboard។", "បង្កើត Walk-in ហើយបង្ហាញ Room Board / Housekeeping។", "បញ្ចប់ដោយ Folio, Payment, Report និងជំហាន MySQL/OTA។"], .82, 5.02, 11.0, 1.30, 11.2, INK, 8, GOLD)
rounded(s, .78, 6.45, 11.84, .40, NAVY, None)
text(s, "Aurelia Collection: បទពិសោធន៍ភ្ញៀវប្រកបដោយភាពប្រណិត ភ្ជាប់ជាមួយប្រព័ន្ធប្រតិបត្តិការដែលអាច Audit បាន។", 1.02, 6.57, 11.35, .14, 9.6, WHITE, True, align=PP_ALIGN.CENTER, margin=0)
footer(s, 16)

prs.core_properties.title = "Aurelia Collection — គម្រោងគ្រប់គ្រងសណ្ឋាគារ · ១៦ ស្លាយជាភាសាខ្មែរ"
prs.core_properties.subject = "បទបង្ហាញខ្មែរខ្លី និងមានខ្លឹមសារ អំពី Website, PMS, Roles, Channels, Pricing និង Deployment"
prs.core_properties.author = "Aurelia Collection"
prs.core_properties.keywords = "Khmer, hotel management, PMS, roles, website, walk-in, OTA, pricing, deployment"
prs.core_properties.comments = "Humanized 16-slide Khmer project deck based on the validated latest release."
prs.save(OUT)
print(OUT)
