from __future__ import annotations

from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
BASE = ROOT / "Aurelia_Collection_Latest_Release_Recap.pptx"
OUT = ROOT / "Aurelia_Collection_Project_Operations_Guide.pptx"

NAVY = RGBColor(11, 20, 39)
NAVY2 = RGBColor(25, 52, 92)
NAVY3 = RGBColor(43, 77, 119)
GOLD = RGBColor(195, 161, 90)
PALE_GOLD = RGBColor(243, 227, 184)
IVORY = RGBColor(250, 248, 244)
BG = RGBColor(247, 248, 250)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(15, 23, 42)
MUTED = RGBColor(100, 116, 139)
GREEN = RGBColor(6, 118, 71)
BLUE = RGBColor(0, 122, 255)
LINE = RGBColor(226, 232, 240)
LIGHT_BLUE = RGBColor(230, 240, 252)
LIGHT_GREEN = RGBColor(236, 253, 243)
LIGHT_AMBER = RGBColor(255, 251, 235)

prs = Presentation(str(BASE))
blank = prs.slide_layouts[6]


def shape(slide, kind, x, y, w, h, fill=None, line=None, transparency=0):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
        shp.fill.transparency = transparency
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.8)
    return shp


def rounded(slide, x, y, w, h, fill=WHITE, line=None, transparency=0):
    return shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, transparency)


def rect(slide, x, y, w, h, fill, line=None, transparency=0):
    return shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, w, h, fill, line, transparency)


def circle(slide, x, y, d, fill, line=None):
    return shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, y, d, d, fill, line)


def text(slide, value, x, y, w, h, size=14, color=INK, bold=False, font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=11, color=INK, gap=5, bullet_color=GOLD):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.text = ""
        r = p.add_run()
        r.text = "•  "
        r.font.name = "Aptos"
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = bullet_color
        r2 = p.add_run()
        r2.text = item
        r2.font.name = "Aptos"
        r2.font.size = Pt(size)
        r2.font.color.rgb = color
    return box


def header(slide, kicker, title, subtitle, num, dark=False):
    title_color = WHITE if dark else INK
    muted = RGBColor(205, 216, 232) if dark else MUTED
    text(slide, kicker.upper(), 0.62, 0.43, 5.2, 0.22, 9.5, GOLD, True, margin=0)
    text(slide, title, 0.62, 0.70, 11.7, 0.58, 27, title_color, True, "Georgia", margin=0)
    text(slide, subtitle, 0.64, 1.32, 11.3, 0.34, 11.5, muted, margin=0)
    text(slide, f"{num:02d}", 12.35, 0.45, 0.4, 0.22, 9, muted, True, align=PP_ALIGN.RIGHT, margin=0)


def footer(slide, num, dark=False):
    color = RGBColor(185, 198, 218) if dark else RGBColor(148, 163, 184)
    rect(slide, 0.62, 7.13, 12.08, 0.008, GOLD if dark else LINE)
    text(slide, "Aurelia Collection · Project operations guide · 27 September 2026", 0.64, 7.19, 8.2, 0.18, 8.2, color, margin=0)
    text(slide, f"{num:02d}", 12.15, 7.16, 0.55, 0.22, 8.5, color, True, align=PP_ALIGN.RIGHT, margin=0)


def add_bg(slide, color=BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def tag(slide, x, y, label, fill=PALE_GOLD, color=NAVY, w=0.72):
    rounded(slide, x, y, w, 0.30, fill, None)
    text(slide, label, x, y + 0.07, w, 0.12, 8.5, color, True, align=PP_ALIGN.CENTER, margin=0)


def card(slide, x, y, w, h, title, desc, fill=WHITE, line=LINE, accent=GOLD, title_color=INK):
    rounded(slide, x, y, w, h, fill, line)
    rect(slide, x, y, 0.055, h, accent)
    text(slide, title, x + 0.22, y + 0.22, w - 0.35, 0.23, 13, title_color, True, "Georgia", margin=0)
    text(slide, desc, x + 0.22, y + 0.60, w - 0.35, h - 0.70, 10.2, MUTED if fill != NAVY else RGBColor(208, 220, 236), margin=0)


# Slide 11 — end-to-end project journey
slide = prs.slides.add_slide(blank)
add_bg(slide, IVORY)
header(slide, "Project journey", "From launch to a settled stay", "The system connects the public guest experience to daily hotel operations and final financial reporting.", 11)
steps = [
    ("01", "Launch", "SQLite starts automatically. Public homepage opens signed out."),
    ("02", "Discover", "Guest explores properties, rooms, photos and maps."),
    ("03", "Price", "Availability, tiered rates, offer and 5% tax are calculated."),
    ("04", "Reserve", "Booking, folio, payment and rate snapshot are created."),
    ("05", "Operate", "Front desk assigns rooms; housekeeping and maintenance work."),
    ("06", "Stay", "Guest is checked in, in-house, then checked out."),
    ("07", "Settle", "Room becomes dirty, task is created and folio is settled."),
    ("08", "Learn", "Managers review revenue, occupancy, channel mix and audit data."),
]
for i, (num, title, desc) in enumerate(steps):
    col = i % 4
    row = i // 4
    x = 0.78 + col * 3.12
    y = 1.95 + row * 2.10
    rounded(slide, x, y, 2.56, 1.55, WHITE, LINE)
    circle(slide, x + 0.22, y + 0.22, 0.38, NAVY if i in (0, 7) else GOLD)
    text(slide, num, x + 0.22, y + 0.315, 0.38, 0.12, 8.5, WHITE if i in (0, 7) else NAVY, True, align=PP_ALIGN.CENTER, margin=0)
    text(slide, title, x + 0.76, y + 0.24, 1.55, 0.22, 13, NAVY, True, "Georgia", margin=0)
    text(slide, desc, x + 0.22, y + 0.78, 2.10, 0.52, 9.7, MUTED, margin=0)
    if i in (3, 7):
        text(slide, "↓", x + 1.15, y + 1.64, 0.28, 0.25, 16, GOLD, True, align=PP_ALIGN.CENTER, margin=0)
    elif col < 3:
        text(slide, "→", x + 2.65, y + 0.64, 0.28, 0.22, 17, GOLD, True, align=PP_ALIGN.CENTER, margin=0)
rounded(slide, 0.78, 6.28, 11.82, 0.42, NAVY, None)
text(slide, "One source of truth: the booking snapshot travels with the guest from reservation to folio to reporting.", 1.02, 6.40, 11.3, 0.16, 10.5, WHITE, True, align=PP_ALIGN.CENTER, margin=0)
footer(slide, 11)

# Slide 12 — channels
slide = prs.slides.add_slide(blank)
add_bg(slide, BG)
header(slide, "Booking channels", "Four ways a reservation can enter the system", "The PMS uses one operational workflow while preserving where each booking came from.", 12)
channels = [
    ("WEBSITE", "Website", "web", "Guest self-service", ["Guest searches and books", "Eligible registered guest receives 10% first-stay offer", "Online guest can view My Stays and cancel"] , NAVY),
    ("WALK-IN", "Walk-in", "walk_in", "Front desk", ["No online account required", "Receptionist enters guest and stay details", "Cash, card, wallet or bank transfer can be recorded"] , GOLD),
    ("PHONE", "Phone", "phone", "Front desk", ["Receptionist books for the caller", "Availability and pricing are validated live", "Guest contact details are stored on the reservation"] , NAVY3),
    ("OTA", "OTA / Agency", "ota", "Manual staff entry today", ["Source is preserved for reporting", "External reference can be kept in notes", "Automatic channel-manager sync is a future step"] , GREEN),
]
for i, (taglabel, title, source, owner, items, accent) in enumerate(channels):
    x = 0.72 + (i % 2) * 6.18
    y = 1.90 + (i // 2) * 2.37
    rounded(slide, x, y, 5.72, 2.05, WHITE, LINE)
    tag(slide, x + 0.26, y + 0.24, taglabel, fill=accent if i in (0, 2) else PALE_GOLD, color=WHITE if i in (0, 2) else NAVY, w=0.98)
    text(slide, title, x + 1.42, y + 0.23, 1.85, 0.23, 14, INK, True, "Georgia", margin=0)
    text(slide, f"Source code: {source}  ·  Owner: {owner}", x + 1.42, y + 0.58, 3.8, 0.18, 9.5, MUTED, margin=0)
    bullets(slide, items, x + 0.30, y + 0.94, 5.05, 0.84, 10, INK, gap=4, bullet_color=accent)
rounded(slide, 0.72, 6.67, 11.9, 0.32, LIGHT_AMBER, RGBColor(253, 230, 138))
text(slide, "Current boundary: OTA/agency reservations are recorded manually; live Booking.com/Expedia-style synchronization is not yet implemented.", 0.98, 6.76, 11.35, 0.14, 9.5, RGBColor(120, 83, 12), True, align=PP_ALIGN.CENTER, margin=0)
footer(slide, 12)

# Slide 13 — roles matrix
slide = prs.slides.add_slide(blank)
add_bg(slide, IVORY)
header(slide, "Roles & access", "Who can do what", "The release has guest access, staff access, and manager/administrator business gates.", 13)
# column headers
x0, y0 = 0.72, 1.88
colx = [x0, 2.52, 6.03, 9.03]
colw = [1.62, 3.34, 2.82, 3.57]
for x, w, label in zip(colx, colw, ["ROLE", "PRIMARY JOB", "CAN ACCESS", "BOUNDARY / NOTES"]):
    rounded(slide, x, y0, w, 0.46, NAVY, None)
    text(slide, label, x + 0.12, y0 + 0.15, w - 0.24, 0.13, 8.7, PALE_GOLD, True, margin=0)
rows = [
    ("Guest", "Book and manage own stay", "Public site · My Stays · self-service cancellation", "Cannot enter staff PMS or view other guests."),
    ("Receptionist", "Front desk and guest movement", "Reservations · calendar · channels · check-in/out · payments", "Operational staff access; no rates, reports or analytics gate."),
    ("Housekeeping", "Room readiness and tasks", "Room Board · housekeeping · maintenance", "Primary job is rooms; current broad staff gate also permits some shared staff pages."),
    ("Manager", "Commercial and property performance", "All staff areas + rates · finance reports · analytics", "Manager-level gate required for rate rules and business reporting."),
    ("Administrator", "System and hotel administration", "Everything + Django admin · users · configuration", "Highest application role; manager privileges included."),
]
for i, row in enumerate(rows):
    y = 2.42 + i * 0.78
    fill = WHITE if i % 2 == 0 else RGBColor(246, 248, 251)
    for x, w in zip(colx, colw):
        rect(slide, x, y, w, 0.74, fill, LINE)
    text(slide, row[0], colx[0] + 0.12, y + 0.20, colw[0] - 0.22, 0.18, 11, NAVY, True, "Georgia", margin=0)
    text(slide, row[1], colx[1] + 0.12, y + 0.14, colw[1] - 0.22, 0.34, 10, INK, True, margin=0)
    text(slide, row[2], colx[2] + 0.12, y + 0.12, colw[2] - 0.22, 0.40, 9.3, MUTED, margin=0)
    text(slide, row[3], colx[3] + 0.12, y + 0.12, colw[3] - 0.22, 0.40, 9.3, MUTED, margin=0)
rounded(slide, 0.72, 6.46, 11.88, 0.48, LIGHT_BLUE, RGBColor(191, 219, 254))
text(slide, "Implementation note: the current RBAC has two major technical gates — staff and manager/admin. Fine-grained housekeeping-only restrictions can be added later if required.", 0.98, 6.60, 11.35, 0.16, 9.5, NAVY3, True, align=PP_ALIGN.CENTER, margin=0)
footer(slide, 13)

# Slide 14 — services by role
slide = prs.slides.add_slide(blank)
add_bg(slide, BG)
header(slide, "Service map", "The right person acts at the right moment", "This is the operational handoff from guest request to hotel team to financial close.", 14)
# vertical service line
line_x = 1.05
rect(slide, line_x + 0.17, 2.16, 0.025, 3.90, GOLD)
service_rows = [
    ("Guest", "Discover and request", "Website, property picker, room search, welcome offer, payment intent and My Stays.", NAVY),
    ("Reception", "Convert and control", "Walk-in, phone and OTA/agency reservations; availability; room assignment; check-in/out.", GOLD),
    ("Housekeeping", "Prepare and restore", "Vacant-clean status, task start/complete, post-checkout room recovery and maintenance reporting.", GREEN),
    ("Finance", "Record and settle", "Folio, invoice, full-stay payment, refund entry and outstanding balance.", BLUE),
    ("Manager", "Review and improve", "Rates, rules, occupancy, ADR, RevPAR, revenue, source mix and property performance.", NAVY3),
]
for i, (role, action, desc, accent) in enumerate(service_rows):
    y = 1.95 + i * 0.88
    circle(slide, 0.87, y + 0.17, 0.42, accent)
    text(slide, str(i + 1), 0.87, y + 0.28, 0.42, 0.12, 9, WHITE, True, align=PP_ALIGN.CENTER, margin=0)
    rounded(slide, 1.52, y, 10.65, 0.68, WHITE, LINE)
    text(slide, role, 1.78, y + 0.13, 1.3, 0.18, 11.5, accent, True, "Georgia", margin=0)
    text(slide, action, 3.15, y + 0.13, 1.65, 0.18, 11, INK, True, margin=0)
    text(slide, desc, 4.92, y + 0.12, 6.85, 0.30, 10, MUTED, margin=0)
rounded(slide, 0.78, 6.48, 11.82, 0.46, NAVY, None)
text(slide, "Guest-facing hospitality services shown on the public site include airport transfers, 24-hour concierge and limousine service.", 1.02, 6.62, 11.3, 0.16, 10, WHITE, True, align=PP_ALIGN.CENTER, margin=0)
footer(slide, 14)

# Slide 15 — current vs next
slide = prs.slides.add_slide(blank)
add_bg(slide, NAVY)
header(slide, "Handoff clarity", "What is live now — and what comes next", "The current release is ready for use; the next work is mostly environment integration and deeper hotel-system maturity.", 15, dark=True)
rounded(slide, 0.78, 1.92, 5.75, 4.75, RGBColor(23, 40, 68), RGBColor(69, 92, 125))
text(slide, "LIVE IN THIS RELEASE", 1.12, 2.28, 3.25, 0.22, 10, PALE_GOLD, True, margin=0)
bullets(slide, [
    "Public guest website and responsive premium interface",
    "10 properties · 24 room types · 369 rooms",
    "Tiered $30–$460 rates, 5% tax only and real one-time 10% offer",
    "Website, walk-in, phone and OTA/agency source handling",
    "Reservations, room board, housekeeping, maintenance and check-in/out",
    "Folios, payments, finance and analytics with snapshot integrity",
    "SQLite automatic launch; optional MySQL/MariaDB and Navicat kit",
    "23/23 tests and 2,160 financial invariants passed",
], 1.12, 2.78, 4.95, 3.15, 10.5, WHITE, 7, bullet_color=GOLD)
rounded(slide, 6.83, 1.92, 5.77, 4.75, PALE_GOLD, None)
text(slide, "RECOMMENDED NEXT", 7.18, 2.28, 3.25, 0.22, 10, NAVY, True, margin=0)
bullets(slide, [
    "Configure and verify the user-owned MySQL/Navicat connection",
    "Add automatic OTA/channel-manager synchronization",
    "Tighten housekeeping/reception permissions if needed",
    "Add folio line items for minibar, spa, restaurant and transport",
    "Add rate plans, group blocks, waitlist and room moves",
    "Add English/KH localization and broader browser testing",
], 7.18, 2.78, 4.95, 2.50, 10.5, NAVY, 10, bullet_color=GREEN)
rounded(slide, 7.18, 5.72, 4.95, 0.55, WHITE, None)
text(slide, "Final environment step: owner-controlled database verification.", 7.37, 5.91, 4.58, 0.15, 9.5, GREEN, True, align=PP_ALIGN.CENTER, margin=0)
footer(slide, 15, dark=True)

# Slide 16 — closing reference
slide = prs.slides.add_slide(blank)
add_bg(slide, IVORY)
header(slide, "Quick reference", "Aurelia Collection in one sentence", "A premium guest booking experience connected to a practical, auditable hotel operating system.", 16)
rounded(slide, 0.82, 2.05, 11.7, 1.40, NAVY, None)
text(slide, "Guest website  →  reservation channel  →  availability & pricing  →  room operations  →  folio & payment  →  management insight", 1.18, 2.45, 10.95, 0.58, 23, WHITE, True, "Georgia", align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
# compact reference cards
card(slide, 0.82, 4.12, 2.75, 1.50, "Guests", "Discover, book, pay, view My Stays and cancel.", WHITE, LINE, GOLD)
card(slide, 3.78, 4.12, 2.75, 1.50, "Front desk", "Control reservations, room movement and guest arrival.", WHITE, LINE, NAVY3)
card(slide, 6.74, 4.12, 2.75, 1.50, "Operations", "Keep rooms clean, sellable, maintained and visible.", WHITE, LINE, GREEN)
card(slide, 9.70, 4.12, 2.82, 1.50, "Management", "Control rates, finance, analytics and system setup.", WHITE, LINE, BLUE)
rounded(slide, 0.82, 6.12, 11.7, 0.54, LIGHT_GREEN, RGBColor(167, 243, 208))
text(slide, "Next action: configure the user-owned MySQL/Navicat connection and verify the live environment.", 1.08, 6.31, 11.15, 0.16, 10.5, GREEN, True, align=PP_ALIGN.CENTER, margin=0)
footer(slide, 16)

prs.core_properties.title = "Aurelia Collection — Project Operations Guide"
prs.core_properties.subject = "Start-to-end project summary, roles, access and booking channels"
prs.core_properties.author = "Aurelia Collection"
prs.core_properties.keywords = "hotel, PMS, roles, access, website, walk-in, OTA, agency, operations"
prs.core_properties.comments = "Detailed guide appended to the validated latest release recap."
prs.save(OUT)
print(OUT)
