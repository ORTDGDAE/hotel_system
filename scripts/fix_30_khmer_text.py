from pathlib import Path
from pptx import Presentation
from pptx.util import Pt

path = Path(__file__).resolve().parent.parent / 'Aurelia_Collection_Hotel_Management_Project_30_Slides_KH.pptx'
prs = Presentation(str(path))
M = {
    'Books and manages a personal stay\nPublic website, My Stays and self-service cancellation': 'កក់ និងគ្រប់គ្រងការស្នាក់នៅរបស់ខ្លួន\nគេហទំព័រសាធារណៈ, My Stays និងលុបចោលដោយខ្លួនឯង',
    'Converts demand into stays\nWalk-in, phone, OTA/agency, arrivals and departures': 'បម្លែងតម្រូវការទៅជាការស្នាក់នៅ\nWalk-in, ទូរស័ព្ទ, OTA/Agency, ភ្ញៀវមកដល់ និងចាកចេញ',
    'Keeps rooms ready to sell\nRoom Board, tasks, room status and maintenance': 'រក្សាបន្ទប់ឱ្យត្រៀមលក់\nRoom Board, កិច្ចការ, ស្ថានភាពបន្ទប់ និង Maintenance',
    'Controls commercial performance\nRates, finance reports, analytics and property scope': 'គ្រប់គ្រងលទ្ធផលពាណិជ្ជកម្ម\nRate, របាយការណ៍ហិរញ្ញវត្ថុ, Analytics និងវិសាលភាពសណ្ឋាគារ',
    'Maintains the platform\nUsers, configuration, Django admin and all manager access': 'ថែទាំ និងគ្រប់គ្រងប្រព័ន្ធ\nអ្នកប្រើ, ការកំណត់រចនាសម្ព័ន្ធ, Django Admin និងសិទ្ធិ Manager ទាំងអស់',
    '•  Reception selects OTA / Agency.\n•  Guest, dates, room and source are recorded.\n•  External reference can be kept in notes.\n•  Availability, pricing and folio stay inside Aurelia.\n•  Channel source remains available for reporting.': '•  Reception ជ្រើស OTA / Agency។\n•  ព័ត៌មានភ្ញៀវ, កាលបរិច្ឆេទ, បន្ទប់ និងប្រភពត្រូវបានកត់ត្រា។\n•  អាចរក្សាលេខយោងខាងក្រៅក្នុង Notes។\n•  Availability, តម្លៃ និង Folio នៅក្នុង Aurelia។\n•  ប្រភព Channel នៅតែមានសម្រាប់របាយការណ៍។',
    '•  Channel-manager API for availability and rates.\n•  Booking push webhooks from OTA partners.\n•  Automatic reservation import and source mapping.\n•  Commission and channel-specific rate plans.\n•  Overbooking protection across connected channels.': '•  API របស់ Channel Manager សម្រាប់ Availability និង Rate។\n•  Webhook ទទួលការកក់ពីដៃគូ OTA។\n•  Import ការកក់ និងផ្គូផ្គងប្រភពដោយស្វ័យប្រវត្តិ។\n•  Commission និង Rate Plan តាម Channel។\n•  ការពារ Overbooking រវាង Channel ដែលបានភ្ជាប់។',
    '•  Entry rooms: $30–$58\n•  Suites and city rooms: $64–$172\n•  Villas and premium stays: $216–$356\n•  Highest current catalog rate: $460\n•  Historical old bookings keep their original totals': '•  បន្ទប់ចាប់ផ្តើម: $30–$58\n•  Suite និងបន្ទប់ទីក្រុង: $64–$172\n•  វីឡា និងការស្នាក់នៅ Premium: $216–$356\n•  តម្លៃខ្ពស់បំផុតបច្ចុប្បន្ន: $460\n•  ការកក់ចាស់ៗរក្សាចំនួនសរុបដើម',
}

def replace(shape, value):
    tf=shape.text_frame
    old=tf.paragraphs[0].runs[0] if tf.paragraphs and tf.paragraphs[0].runs else None
    align=tf.paragraphs[0].alignment if tf.paragraphs else None
    size=old.font.size.pt if old and old.font.size else 10
    bold=bool(old.font.bold) if old else False
    color=None
    if old:
        try: color=old.font.color.rgb
        except Exception: pass
    if len(value)>180: size=min(size,8.6)
    elif len(value)>100: size=min(size,9.2)
    tf.clear(); tf.word_wrap=True; p=tf.paragraphs[0]
    if align is not None: p.alignment=align
    r=p.add_run(); r.text=value; r.font.name='Noto Sans Khmer'; r.font.size=Pt(size); r.font.bold=bold
    if color is not None: r.font.color.rgb=color

count=0
for idx, slide in enumerate(prs.slides):
    if idx < 22:
        continue
    for shape in slide.shapes:
        if hasattr(shape,'text_frame') and shape.text in M:
            replace(shape,M[shape.text]); count+=1
prs.save(path)
print('fixed',count,'combined Khmer text boxes')
