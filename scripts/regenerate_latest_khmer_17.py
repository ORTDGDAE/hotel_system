from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "Aurelia_Collection_Hotel_Management_Project_17_Slides_KH_Latest.pptx"
prs = Presentation(str(OUT))

# Preserve the existing designed 17-slide deck while refreshing the copy with
# the current multi-property, property-manager, housekeeping-team and QA facts.
REPLACEMENTS = {
    "គម្រោងគ្រប់គ្រង សណ្ឋាគារ": "គម្រោងគ្រប់គ្រងសណ្ឋាគារ ពហុទីតាំង",
    "ទិដ្ឋភាពសង្ខេប និងមានខ្លឹមសារ អំពីដំណើររបស់ភ្ញៀវទាំងមូល, ប្រតិបត្តិការ PMS, សិទ្ធិតាមតួនាទី, តម្លៃ និងផ្លូវដាក់ឱ្យប្រើ។": "ទិដ្ឋភាពសង្ខេបអំពីដំណើរភ្ញៀវ, PMS, សិទ្ធិតាមសណ្ឋាគារ, តម្លៃ និងផ្លូវដាក់ឱ្យប្រើ។",
    "ការកក់ផ្ទាល់ · Front Desk · Housekeeping · ហិរញ្ញវត្ថុ · Analytics": "10 Hotels · 10 Property Managers · 30 Housekeepers",
    "សង្ខេបគម្រោងកំណែចុងក្រោយ · ២៩ កញ្ញា ២០២៦": "សង្ខេបគម្រោងកំណែចុងក្រោយ · ០៤ តុលា ២០២៦",
    "Aurelia Collection · គម្រោងគ្រប់គ្រងសណ្ឋាគារ · ២៩ កញ្ញា ២០២៦": "Aurelia Collection · គម្រោងគ្រប់គ្រងសណ្ឋាគារ · ០៤ តុលា ២០២៦",
    "លំហូរប្រតិបត្តិការដែលច្បាស់: ស្វែងរក, កក់, រៀបចំ, មកដល់, ស្នាក់នៅ, Check-out, ផ្ទៀងផ្ទាត់ និងកែលម្អ។": "លំហូរប្រតិបត្តិការដែលច្បាស់: ស្វែងរក, កក់, Check-in, ស្នាក់នៅ, Payment, Check-out និង Housekeeping។",
    "ប្រព័ន្ធប្រតិបត្តិការតែមួយ ភ្ជាប់តម្រូវការរបស់ភ្ញៀវ ជាមួយការរៀបចំបន្ទប់, ការគ្រប់គ្រងហិរញ្ញវត្ថុ និងព័ត៌មានសម្រាប់អ្នកគ្រប់គ្រង។": "ប្រព័ន្ធតែមួយ ភ្ជាប់ភ្ញៀវ, Reservation, Room Board, Housekeeping, Finance និងព័ត៌មានសម្រាប់អ្នកគ្រប់គ្រងទាំង ១០ សណ្ឋាគារ។",
    "5%": "5%",
    "តួនាទី ៥ ធ្វើឱ្យសណ្ឋាគារដំណើរការ": "តួនាទី ៦ និងវិសាលភាពតាមសណ្ឋាគារ",
    "តួនាទីនីមួយៗឃើញការងារដែលខ្លួនទទួលខុសត្រូវ ខណៈ Platform ការពារស្តុករួម និងកំណត់ត្រាហិរញ្ញវត្ថុ។": "Property Manager គ្រប់គ្រងសណ្ឋាគាររបស់ខ្លួន ខណៈ Group Manager មើលគ្រប់ ១០ សណ្ឋាគារ និង Platform ការពារស្តុក និងហិរញ្ញវត្ថុ។",
    "Manager": "Property Manager",
    "Rate, ហិរញ្ញវត្ថុ, Analytics និងវិសាលភាពសណ្ឋាគារ": "Reservation, Room, Housekeeping, Finance និង Analytics របស់សណ្ឋាគារខ្លួន",
    "Administrator": "Group Manager / Admin",
    "អ្នកប្រើ, ការកំណត់រចនា និងការគ្រប់គ្រងប្រព័ន្ធពេញលេញ": "Group Manager មើលគ្រប់សណ្ឋាគារ; Admin គ្រប់គ្រង Users និងប្រព័ន្ធពេញលេញ",
    "គោលការណ៍សិទ្ធិ": "គោលការណ៍សិទ្ធិ និងវិសាលភាព",
    "ភ្ញៀវឃើញដំណើររបស់ខ្លួន · Staff ប្រតិបត្តិការសណ្ឋាគារ · Manager គ្រប់គ្រងអាជីវកម្ម · Administrator គ្រប់គ្រងប្រព័ន្ធ។": "ភ្ញៀវឃើញដំណើររបស់ខ្លួន · Reception ប្រតិបត្តិការក្នុងសណ្ឋាគារ · Property Manager គ្រប់គ្រងមួយសណ្ឋាគារ · Group Manager មើលទាំងក្រុម។",
    "Housekeeping និងស្ថានភាពបន្ទប់ ធានាថាបន្ទប់អាចដាក់លក់បាន។": "ឃើញតែកិច្ចការដែលបានប្រគល់ និងធ្វើឱ្យបន្ទប់ត្រឡប់ទៅ Clean។",
    "Check-out": "Check-out",
    "Folio, ការទូទាត់, ស្ថានភាព និង Audit Trail ត្រូវបានបិទបញ្ចប់។": "Folio, Deposit/Payment, ស្ថានភាព Dirty និង Audit Trail ត្រូវបានបិទបញ្ចប់។",
    "តម្លៃប្រតិបត្តិការ: Reception, Housekeeping និង Management ឃើញទិន្នន័យពិតដូចគ្នា ក្នុងពេលតែមួយ។": "តម្លៃប្រតិបត្តិការ: Property Manager ប្រគល់ Task, Housekeeper ធ្វើការងារ និង Group Manager មើលលទ្ធផលទាំងក្រុម។",
    "រៀបចំ": "រៀបចំបន្ទប់",
    "Housekeeping និងស្ថានភាពបន្ទប់ ធានាថាបន្ទប់អាចដាក់លក់បាន។": "Room Board បង្ហាញ Dirty ហើយបង្កើត Cleaning Task សម្រាប់ Property Manager ប្រគល់ទៅក្រុម Housekeeping។",
    "Check-out, Folio និងការទូទាត់។": "Check-out, Folio, Deposit/Payment និង Cleaning Task។",
    "បិទការស្នាក់នៅ": "បិទការស្នាក់នៅ និងទូទាត់",
    "បង្ហាញ Check-out, Folio និងការទូទាត់។": "បង្ហាញ Check-out, Folio, Deposit និងការប្រគល់ Cleaning Task។",
    "បញ្ចូលព័ត៌មានទំនាក់ទំនង, សំណើ និងវិធីទូទាត់។": "បញ្ចូលព័ត៌មានទំនាក់ទំនង, Phone number, សំណើ និងវិធីទូទាត់។",
    "ទំនុកចិត្តលើកំណែចេញផ្សាយ": "ទំនុកចិត្តលើកំណែចេញផ្សាយ",
    "តេស្ត ២៣/២៣ បានជោគជ័យ · Ledger Invariant ២,១៦០ បានជោគជ័យ · បានធ្វើបច្ចុប្បន្នភាពសណ្ឋាគារ ១០ និងប្រភេទបន្ទប់ ២៤": "តេស្ត ៤១/៤១ ជោគជ័យ · Ledger Invariant ២,១៦០ ជោគជ័យ · Property Manager ១០ · Housekeeper ៣០",
    "សម្រាប់ក្រុមការងារ": "សម្រាប់ក្រុមការងារ",
    "Reception, Housekeeping និង Manager ប្រើទិន្នន័យតែមួយ ដើម្បីគ្រប់គ្រងការកក់, បន្ទប់ និងការស្នាក់នៅ។": "Reception, Property Manager និង Housekeeping ប្រើទិន្នន័យតែមួយ; Group Manager មើលលទ្ធផលគ្រប់ ១០ សណ្ឋាគារ។",
    "ម្ចាស់សណ្ឋាគារឃើញ Occupancy, Revenue, Payment និង Folio ដោយមាន Snapshot និង Audit ដែលអាចទុកចិត្តបាន។": "ម្ចាស់សណ្ឋាគារឃើញ Occupancy, Revenue, Payment, Folio និងក្រុមការងារ ដោយមាន Snapshot និង Audit ដែលអាចទុកចិត្តបាន។",
    "Aurelia Collection: បទពិសោធន៍ភ្ញៀវប្រកបដោយភាពប្រណិត ភ្ជាប់ជាមួយប្រព័ន្ធប្រតិបត្តិការដែលអាច Audit បាន។": "Aurelia Collection: បទពិសោធន៍ភ្ញៀវប្រកបដោយភាពប្រណិត ភ្ជាប់ជាមួយប្រព័ន្ធប្រតិបត្តិការតាមសណ្ឋាគារ ដែលអាច Audit បាន។",
    "គោលការណ៍ Snapshot": "Snapshot និង Deposit",
    "Folio និង Payment អាចផ្ទៀងផ្ទាត់ត្រឡប់ទៅ Booking។": "Folio, Deposit និង Payment អាចផ្ទៀងផ្ទាត់ត្រឡប់ទៅ Booking; បង់លើសចំនួនត្រូវបានបដិសេធ។",
    "Service Charge: 0% · Tax: 5%": "Service Charge: 0% · Tax: 5% · Deposit អាចកត់ត្រាបាន",
    "កំណែបច្ចុប្បន្នត្រៀមសម្រាប់ Demo និងប្រើប្រាស់; ជំហានបន្ទាប់គឺការភ្ជាប់បរិស្ថានរបស់ម្ចាស់ប្រព័ន្ធ។": "កំណែបច្ចុប្បន្នត្រៀមសម្រាប់ Demo និងប្រើប្រាស់; ឥឡូវមាន ១០ Property Manager និង ៣០ Housekeepers បែងចែកតាមសណ្ឋាគារ។",
    "បង្កើត Walk-in ហើយបង្ហាញ Room Board / Housekeeping។": "បង្កើត Walk-in ហើយបង្ហាញ Room Board, Dirty Task និងក្រុម Housekeeping។",
    "បញ្ចប់ដោយ Folio, Payment, Report និងជំហាន MySQL/OTA។": "បញ្ចប់ដោយ Folio, Deposit/Payment, Report និងជំហាន MySQL/OTA។",
    "ប្រព័ន្ធតែមួយ សម្រាប់ភ្ញៀវ ក្រុមសណ្ឋាគារ និងម្ចាស់អាជីវកម្ម — ចាប់ពីការស្វែងរក រហូតដល់ការទូទាត់ និងរបាយការណ៍។": "ប្រព័ន្ធតែមួយ សម្រាប់ភ្ញៀវ ក្រុមសណ្ឋាគារ និងម្ចាស់អាជីវកម្ម — ១០ សណ្ឋាគារ ចាប់ពីការស្វែងរក រហូតដល់ការទូទាត់ និងរបាយការណ៍។",
    "ស្វែងរកបន្ទប់, កក់ដោយងាយ, មើលតម្លៃសរុប, គ្រប់គ្រង My Stays និងទទួលសេវាដែលច្បាស់លាស់។": "ស្វែងរកបន្ទប់, កក់ដោយងាយ, មើលតម្លៃសរុប, បញ្ចូល Phone number និងគ្រប់គ្រង My Stays។",
    "កំណែនេះត្រៀមសម្រាប់ការបង្ហាញ និងការផ្ទៀងផ្ទាត់ Database របស់ម្ចាស់ប្រព័ន្ធ។ ជំហានបន្ទាប់: MySQL/Navicat និង OTA Channel Manager។": "កំណែនេះត្រៀមសម្រាប់ការបង្ហាញ និងការផ្ទៀងផ្ទាត់ Database របស់ម្ចាស់ប្រព័ន្ធ។ ស្ថានភាព QA: ៤១/៤១ តេស្ត និង ២,១៦០ Ledger checks បានជោគជ័យ។",
    "ប្រភេទបន្ទប់ ២៤": "ប្រភេទ ២៤ · បន្ទប់ ៣៦៩",
    "ទិន្នន័យបន្ទប់": "បន្ទប់ក្នុងប្រព័ន្ធ",
    "សណ្ឋាគារ ១០": "សណ្ឋាគារ ១០",
}


def replace_in_shape(shape):
    if not hasattr(shape, "text_frame"):
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            value = run.text
            for old, new in REPLACEMENTS.items():
                value = value.replace(old, new)
            run.text = value

for slide in prs.slides:
    for shape in slide.shapes:
        replace_in_shape(shape)

# Clean up replacement-order collisions (e.g. "Manager" inside an already
# expanded "Property Manager" phrase).
CLEANUP = {
    "Property Property Manager": "Property Manager",
    "Group Property Manager": "Group Manager",
    "Channel Property Manager": "Channel Manager",
}
for slide in prs.slides:
    for shape in slide.shapes:
        if not hasattr(shape, "text_frame"):
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                value = run.text
                for old, new in CLEANUP.items():
                    while old in value:
                        value = value.replace(old, new)
                run.text = value

# Slide 14 is the role matrix. Make the manager column explicit without
# changing the designed table geometry.
slide14 = prs.slides[13]
for shape in slide14.shapes:
    if hasattr(shape, "text") and shape.text.strip() == "MANAGER":
        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                run.text = "PROPERTY / GROUP"
    if hasattr(shape, "text") and shape.text.strip() == "គោលការណ៍: អ្នកប្រើឃើញតែការងារដែលខ្លួនទទួលខុសត្រូវ ហើយប្រព័ន្ធការពារស្តុក និងហិរញ្ញវត្ថុរួម។":
        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                run.text = "គោលការណ៍: Property Manager គ្រប់គ្រងតែសណ្ឋាគារខ្លួន; Group Manager មើលគ្រប់សណ្ឋាគារ; Housekeeper ឃើញតែកិច្ចការដែលបានប្រគល់។"

# Final whole-shape cleanup for phrases that may cross PowerPoint run
# boundaries. Only affected text boxes are rewritten.
FINAL_CLEANUP = {
    "Property Property Manager": "Property Manager",
    "Group Property Manager": "Group Manager",
    "Channel Property Manager": "Channel Manager",
    "គោលការណ៍សិទ្ធិ និងវិសាលភាព និងវិសាលភាព": "គោលការណ៍សិទ្ធិ និងវិសាលភាព",
    "គម្រោងគ្រប់គ្រង\nសណ្ឋាគារ": "គម្រោងគ្រប់គ្រងសណ្ឋាគារ ពហុទីតាំង",
    "កំណែនេះត្រៀមសម្រាប់ការបង្ហាញ និងការផ្ទៀងផ្ទាត់ Database របស់ម្ចាស់ប្រព័ន្ធ។ ជំហានបន្ទាប់: MySQL/Navicat និង OTA Channel Manager។": "កំណែនេះត្រៀមសម្រាប់ការបង្ហាញ និងការផ្ទៀងផ្ទាត់ Database របស់ម្ចាស់ប្រព័ន្ធ។ ស្ថានភាព QA: ៤១/៤១ តេស្ត និង ២,១៦០ Ledger checks បានជោគជ័យ។",
    "Reception, Housekeeping និង Property Manager ប្រើទិន្នន័យតែមួយ ដើម្បីគ្រប់គ្រងការកក់, បន្ទប់ និងការស្នាក់នៅ។": "Reception, Property Manager និង Housekeeping ប្រើទិន្នន័យតែមួយ; Group Manager មើលលទ្ធផលគ្រប់ ១០ សណ្ឋាគារ។",
    "ភ្ញៀវឃើញដំណើររបស់ខ្លួន · Staff ប្រតិបត្តិការសណ្ឋាគារ · Property Manager គ្រប់គ្រងអាជីវកម្ម · Group Manager / Admin គ្រប់គ្រងប្រព័ន្ធ។": "ភ្ញៀវឃើញដំណើររបស់ខ្លួន · Reception ប្រតិបត្តិការក្នុងសណ្ឋាគារ · Property Manager គ្រប់គ្រងមួយសណ្ឋាគារ · Group Manager មើលទាំងក្រុម។",
    "រៀបចំបន្ទប់បន្ទប់": "រៀបចំបន្ទប់",
    "បានរៀបចំបន្ទប់": "បានរៀបចំ",
    "Deposit អាចកត់ត្រាបាន · Deposit អាចកត់ត្រាបាន": "Deposit អាចកត់ត្រាបាន",
    "បិទការស្នាក់នៅ និងទូទាត់ និងទូទាត់": "បិទការស្នាក់នៅ និងទូទាត់",
    "តេស្ត ៤១/៤១ ជោគជ័យ · Ledger Invariant ២,១៦០ ជោគជ័យ · Property Manager ១០ · Housekeeper ៣០": "តេស្ត ៤១/៤១ ជោគជ័យ · Ledger Invariant ២,១៦០ ជោគជ័យ · Property Manager ១០ · Housekeeper ៣០ (៣ នាក់ក្នុងមួយសណ្ឋាគារ)",
    "កំណែបច្ចុប្បន្នត្រៀមសម្រាប់ Demo និងប្រើប្រាស់; ឥឡូវមាន ១០ Property Manager និង ៣០ Housekeepers បែងចែកតាមសណ្ឋាគារ។": "កំណែបច្ចុប្បន្នត្រៀមសម្រាប់ Demo និងប្រើប្រាស់; ឥឡូវមាន ១០ Property Manager (១ ក្នុងមួយសណ្ឋាគារ) និង ៣០ Housekeepers (៣ ក្នុងមួយសណ្ឋាគារ)។",
    "Room Board បង្ហាញ Dirty ហើយបង្កើត Cleaning Task សម្រាប់ Property Manager ប្រគល់ទៅក្រុម Housekeeping។": "Room Board បង្ហាញ Dirty ហើយបង្កើត Cleaning Task សម្រាប់ Property Manager ប្រគល់ទៅក្រុម Housekeeping ដោយបង្ហាញលេខបន្ទប់ និងប្រភេទបន្ទប់។",
    "Folio, Deposit និង Payment អាចផ្ទៀងផ្ទាត់ត្រឡប់ទៅ Booking; បង់លើសចំនួនត្រូវបានបដិសេធ។": "Folio, Deposit និង Payment អាចផ្ទៀងផ្ទាត់ត្រឡប់ទៅ Booking; Partial Payment អាចកត់ត្រាតាមដំណាក់កាល, បង់លើសចំនួនត្រូវបានបដិសេធ និង Check-out ត្រូវទូទាត់សមតុល្យ។",
    "ការកែ Catalog ថ្មី មិនកែចំណូលប្រវត្តិសាស្ត្រ។": "ការកែ Catalog ថ្មី មិនកែចំណូលប្រវត្តិសាស្ត្រ; Invoice ចាស់ និង Pricing Snapshot ចាស់ មិនត្រូវបានសរសេរឡើងវិញ។",
    "បច្ចេកវិទ្យាអាចពង្រីកបាន ដោយមិនប្តូរភាសាប្រតិបត្តិការរបស់សណ្ឋាគារ។": "បច្ចេកវិទ្យាអាចពង្រីកបាន ដោយមិនប្តូរភាសាប្រតិបត្តិការរបស់សណ្ឋាគារ។ • Metrics បច្ចុប្បន្ន: ៣៦៩ បន្ទប់ · ប្រភេទ ២៤ · Booking ១៧០ · Folio ១៧០ · Payment ១០៣។",
}
for slide in prs.slides:
    for shape in slide.shapes:
        if not hasattr(shape, "text"):
            continue
        value = shape.text
        fixed = value
        for old, new in FINAL_CLEANUP.items():
            # Some enrichment replacements intentionally retain the old
            # phrase as a prefix; apply those once to avoid a replacement
            # loop, while repeated-artifact cleanups can safely converge.
            if old in new:
                if new not in fixed:
                    fixed = fixed.replace(old, new)
            else:
                while old in fixed:
                    fixed = fixed.replace(old, new)
        if fixed != value:
            shape.text = fixed
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Noto Sans Khmer"

prs.core_properties.title = "Aurelia Collection — គម្រោងគ្រប់គ្រងសណ្ឋាគារ · ១៧ ស្លាយជាភាសាខ្មែរ · កំណែ ០៤ តុលា ២០២៦"
prs.core_properties.subject = "Khmer latest project deck with multi-property managers, 30 housekeeping cleaners, deposits, dirty-room workflow and 41-test QA update"
prs.core_properties.comments = "Regenerated from the validated current Aurelia Collection PMS implementation."
prs.save(OUT)
print(OUT)
