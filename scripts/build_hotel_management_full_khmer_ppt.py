from __future__ import annotations

from pathlib import Path
from pptx import Presentation
from pptx.util import Pt

ROOT = Path(__file__).resolve().parent.parent
BASE = ROOT / "Aurelia_Collection_Hotel_Management_Project_Full_Deck.pptx"
OUT = ROOT / "Aurelia_Collection_Hotel_Management_Project_Full_Deck_KH.pptx"
prs = Presentation(str(BASE))

# Natural, presentation-ready Khmer translations. Technical product names and
# file names remain in Latin script where that is clearer for hotel teams.
T = {
    # Shared / cover / recap slides
    "Hotel Booking & Property Management System": "ប្រព័ន្ធកក់សណ្ឋាគារ និងគ្រប់គ្រងប្រតិបត្តិការ",
    "Latest release recap": "សង្ខេបកំណែចុងក្រោយ",
    "A polished guest experience, operational PMS, financially consistent booking flow, and MySQL/Navicat-ready deployment path.": "បទពិសោធន៍ភ្ញៀវប្រកបដោយភាពប្រណិត, PMS សម្រាប់ប្រតិបត្តិការសណ្ឋាគារ, លំហូរកក់ដែលត្រឹមត្រូវផ្នែកហិរញ្ញវត្ថុ និងផ្លូវដាក់ឱ្យប្រើជាមួយ MySQL/Navicat។",
    "BUILD 27 SEPTEMBER 2026": "បង្កើតនៅថ្ងៃទី ២៧ កញ្ញា ២០២៦",
    "Cambodia · 10 sanctuaries · One standard of grace": "កម្ពុជា · សណ្ឋាគារ ១០ កន្លែង · ស្តង់ដារសេវាកម្មតែមួយ",
    "RELEASE SNAPSHOT": "សង្ខេបការចេញផ្សាយ",
    "What shipped in the latest version": "អ្វីដែលបានដាក់ឱ្យប្រើក្នុងកំណែចុងក្រោយ",
    "The product is now a coherent end-to-end hotel platform, not just a booking front end.": "ផលិតផលនេះឥឡូវជាប្រព័ន្ធសណ្ឋាគារពេញលេញពីដើមដល់ចប់ មិនមែនមានតែផ្នែកកក់ប៉ុណ្ណោះទេ។",
    "Cambodia properties": "អចលនទ្រព្យសណ្ឋាគារនៅកម្ពុជា",
    "Room types": "ប្រភេទបន្ទប់",
    "Physical rooms": "បន្ទប់ជាក់ស្តែង",
    "Django tests passing": "តេស្ត Django ដែលបានជោគជ័យ",
    "The release pillars": "សសរស្តម្ភសំខាន់ៗនៃកំណែនេះ",
    "Guest journey": "ដំណើររបស់ភ្ញៀវ",
    "Collection discovery, room search, mobile navigation and signup conversion.": "ស្វែងរកសណ្ឋាគារ, ស្វែងរកបន្ទប់, ប្រើលើទូរស័ព្ទ និងបម្លែងអ្នកចូលមើលទៅជាអ្នកកក់។",
    "PMS operations": "ប្រតិបត្តិការ PMS",
    "Reservations, room board, housekeeping, folios, payments and reporting.": "ការកក់, Room Board, housekeeping, folio, ការទូទាត់ និងរបាយការណ៍។",
    "Commercial logic": "តក្កវិជ្ជាពាណិជ្ជកម្ម",
    "Tiered rack rates, rate-rule records, snapshots and one-time welcome offer.": "តម្លៃបន្ទប់តាមកម្រិត, កំណត់ត្រាច្បាប់តម្លៃ, snapshot និងការផ្តល់ជូនស្វាគមន៍ប្រើបានម្តង។",
    "Deployment path": "ផ្លូវដាក់ឱ្យប្រើ",
    "MySQL/MariaDB configuration, SQLite transfer, Navicat SQL export and one-click scripts.": "ការកំណត់ MySQL/MariaDB, ការផ្ទេរ SQLite, ការនាំចេញ SQL ទៅ Navicat និងស្គ្រីបចុចម្តង។",
    "Release signal": "សញ្ញានៃការចេញផ្សាយ",
    "A premium guest shell backed by operationally credible hotel data and a verified financial ledger.": "ផ្នែកភ្ញៀវដែលមានរូបរាងប្រណិត គាំទ្រដោយទិន្នន័យសណ្ឋាគារដែលអាចប្រើការបាន និងបញ្ជីហិរញ្ញវត្ថុដែលបានផ្ទៀងផ្ទាត់។",
    "Current preview remains available on port 8000 while the user-owned MySQL connection is configured.": "Preview បច្ចុប្បន្ននៅតែដំណើរការតាម port 8000 ខណៈពេលកំពុងរៀបចំការភ្ជាប់ MySQL របស់ម្ចាស់ប្រព័ន្ធ។",
    "GUEST EXPERIENCE": "បទពិសោធន៍ភ្ញៀវ",
    "A calmer, more premium path to booking": "ផ្លូវកក់ដែលសាមញ្ញ ស្ងប់ស្ងាត់ និងមានភាពប្រណិតជាងមុន",
    "The public site now balances editorial hospitality, responsive interaction, and clear conversion moments.": "គេហទំព័រសាធារណៈឥឡូវមានតុល្យភាពរវាងបរិយាកាសបដិសណ្ឋារកិច្ច, ការឆ្លើយតបលើគ្រប់ឧបករណ៍ និងចំណុចជំរុញការកក់ដែលច្បាស់លាស់។",
    "Discover the collection": "ស្វែងរកសណ្ឋាគារក្នុង Collection",
    "Real Cambodia properties, city context, vibe tags, pricing and map links in one glance.": "មើលសណ្ឋាគារពិតនៅកម្ពុជា, ទីតាំងទីក្រុង, ស្លាកបរិយាកាស, តម្លៃ និងតំណផែនទីក្នុងមួយក្រឡា។",
    "Choose with confidence": "ជ្រើសរើសដោយទំនុកចិត្ត",
    "Curated hotel recommendations, animated selection feedback and accessible picker sheets.": "ការណែនាំសណ្ឋាគារដែលបានជ្រើសរើសយ៉ាងយកចិត្តទុកដាក់, ប្រតិកម្មពេលជ្រើសរើស និងបង្អួចជ្រើសរើសដែលងាយប្រើ។",
    "Book the right stay": "កក់ការស្នាក់នៅដែលសមរម្យ",
    "Rooms, suites and island villas lead into a consistent, responsive booking flow.": "បន្ទប់, Suite និងវីឡាកោះ នាំទៅកាន់លំហូរកក់ដែលមានស្តង់ដារតែមួយ និងឆ្លើយតបល្អ។",
    "Guest-facing details": "ព័ត៌មានសម្រាប់ភ្ញៀវ",
    "Responsive navigation · curated hotel picker · Google Maps public links · bfcache-safe page transitions · designed 404 and auth boundaries": "ម៉ឺនុយឆ្លើយតបល្អ · កម្មវិធីជ្រើសសណ្ឋាគារដែលបានរៀបចំ · តំណ Google Maps សាធារណៈ · ការប្តូរទំព័រដែលមានសុវត្ថិភាព · ទំព័រ 404 និងព្រំដែន Login ដែលបានរចនា",
    "COMMERCIAL LOGIC": "តក្កវិជ្ជាពាណិជ្ជកម្ម",
    "Tiered pricing with a real welcome conversion offer": "តម្លៃតាមកម្រិត ជាមួយការផ្តល់ជូនស្វាគមន៍ដែលប្រើបានពិត",
    "Rack rates vary by room and property tier; booking snapshots remain financially authoritative.": "តម្លៃមូលដ្ឋានខុសគ្នាតាមប្រភេទបន្ទប់ និងកម្រិតសណ្ឋាគារ; snapshot នៃការកក់នៅតែជាតម្លៃយោងផ្លូវការផ្នែកហិរញ្ញវត្ថុ។",
    "Current catalog ladder": "ជណ្តើរតម្លៃក្នុង Catalog បច្ចុប្បន្ន",
    "entry rate": "តម្លៃចាប់ផ្តើម",
    "Entry": "កម្រិតចាប់ផ្តើម",
    "Suites": "Suite",
    "Landmark": "សណ្ឋាគារលេចធ្លោ",
    "Private island": "កោះឯកជន",
    "24 room types · rules remain in the PMS but are disabled for flat dated demo searches.": "មានបន្ទប់ ២៤ ប្រភេទ · ច្បាប់តម្លៃនៅតែមានក្នុង PMS ប៉ុន្តែត្រូវបានបិទសម្រាប់ការស្វែងរក Demo តាមកាលបរិច្ឆេទធម្មតា។",
    "10% welcome offer": "ការផ្តល់ជូនស្វាគមន៍ ១០%",
    "Newly registered guests see a delayed modal with a blurred background and unlock the real first-stay discount.": "ភ្ញៀវដែលចុះឈ្មោះថ្មីនឹងឃើញ Modal បន្ទាប់ពីបន្តិច ជាមួយផ្ទៃខាងក្រោយព្រិល ហើយអាចបើកការបញ្ចុះតម្លៃសម្រាប់ការស្នាក់នៅលើកដំបូង។",
    "% OFF": "% បញ្ចុះ",
    "Welcome to Aurelia": "សូមស្វាគមន៍មកកាន់ Aurelia",
    "Register now and save 10% on your first web booking.": "ចុះឈ្មោះឥឡូវនេះ ដើម្បីសន្សំ ១០% សម្រាប់ការកក់តាមគេហទំព័រលើកដំបូង។",
    "Unlock offer  ↗": "បើកការផ្តល់ជូន  ↗",
    "Nightly rates → 10% discount → service → tax → full-stay total": "តម្លៃក្នុងមួយយប់ → បញ្ចុះ ១០% → សេវា → ពន្ធ → សរុបពេញការស្នាក់នៅ",
    "Nightly rates → 10% discount → 5% tax only → full-stay total": "តម្លៃក្នុងមួយយប់ → បញ្ចុះ ១០% → ពន្ធតែ ៥% → សរុបពេញការស្នាក់នៅ",
    "PROPERTY OPERATIONS": "ប្រតិបត្តិការសណ្ឋាគារ",
    "The guest promise is connected to the PMS": "ការសន្យាចំពោះភ្ញៀវ ត្រូវបានភ្ជាប់ជាមួយ PMS",
    "Front desk, operations and finance share the same booking lifecycle and room inventory.": "Front desk, ប្រតិបត្តិការ និងហិរញ្ញវត្ថុ ប្រើវដ្តការកក់ និងស្តុកបន្ទប់តែមួយ។",
    "Reserve": "កក់",
    "Booking + snapshot": "ការកក់ + Snapshot",
    "Assign": "ចាត់បន្ទប់",
    "Room inventory": "ស្តុកបន្ទប់",
    "Operate": "ប្រតិបត្តិការ",
    "Housekeeping + room board": "Housekeeping + Room Board",
    "Settle": "បិទបញ្ជី",
    "Folio + payment": "Folio + ការទូទាត់",
    "Report": "រាយការណ៍",
    "Analytics + audit": "Analytics + Audit",
    "PMS surfaces": "ផ្នែកសំខាន់ៗក្នុង PMS",
    "•  Reservations, calendar and new booking workflow\n•  Room Board with SSE live updates across 369 rooms\n•  Housekeeping tasks, maintenance requests and role gates\n•  Manager, reception, housekeeping, guest and admin shells": "•  ការកក់, ប្រតិទិន និងលំហូរកក់ថ្មី\n•  Room Board ដែលធ្វើបច្ចុប្បន្នភាពផ្ទាល់សម្រាប់បន្ទប់ ៣៦៩\n•  កិច្ចការ Housekeeping, សំណើជួសជុល និងការគ្រប់គ្រងតាមតួនាទី\n•  ផ្ទាំងសម្រាប់ Manager, Reception, Housekeeping, Guest និង Admin",
    "Operational confidence": "ទំនុកចិត្តលើប្រតិបត្តិការ",
    "Every booking carries its own commercial snapshot — rates, discount, service, tax and total — so the PMS never needs to rewrite history.": "ការកក់នីមួយៗមាន Snapshot ពាណិជ្ជកម្មផ្ទាល់ខ្លួន — តម្លៃ, បញ្ចុះតម្លៃ, សេវា, ពន្ធ និងសរុប — ដូច្នេះ PMS មិនចាំបាច់កែប្រវត្តិចាស់ឡើយ។",
    "Full-stay payment totals remain full-stay totals.": "ចំនួនទឹកប្រាក់ទូទាត់ពេញការស្នាក់នៅ នៅតែជាចំនួនពេញការស្នាក់នៅ។",
    "FINANCIAL INTEGRITY": "ភាពត្រឹមត្រូវផ្នែកហិរញ្ញវត្ថុ",
    "The numbers are treated as contracts": "លេខទាំងនេះត្រូវបានចាត់ទុកជាកិច្ចសន្យា",
    "Historical reservations, folios, payments and cancellation bases remain stable after catalog changes.": "ការកក់, Folio, ការទូទាត់ និងមូលដ្ឋានគណនាការលុបចោលចាស់ៗ នៅតែរក្សាដដែល បន្ទាប់ពីកែ Catalog។",
    "Bookings": "ការកក់",
    "Folios": "Folio",
    "Payments": "ការទូទាត់",
    "Invariant checks": "ការត្រួតពិនិត្យ Invariant",
    "Snapshot principle": "គោលការណ៍ Snapshot",
    "A booking stores the rate details and totals that were actually sold. Later price-catalog changes do not mutate historical revenue.": "ការកក់រក្សាទុកព័ត៌មានតម្លៃ និងចំនួនសរុបដែលបានលក់ពិតប្រាកដ។ ការកែតម្លៃ Catalog នៅពេលក្រោយ មិនកែប្រែចំណូលប្រវត្តិសាស្ត្រឡើយ។",
    "The 10% offer is applied before service charge and tax, then stored in the booking snapshot.": "ការបញ្ចុះ ១០% ត្រូវបានអនុវត្តមុនគិតសេវា និងពន្ធ បន្ទាប់មករក្សាទុកក្នុង Snapshot នៃការកក់។",
    "The 10% offer is applied before the 5% tax, then stored in the booking snapshot.": "ការបញ្ចុះ ១០% ត្រូវបានអនុវត្តមុនគិតពន្ធ ៥% បន្ទាប់មករក្សាទុកក្នុង Snapshot នៃការកក់។",
    "Audit result": "លទ្ធផល Audit",
    "PASSED": "ជោគជ័យ",
    "BRAND SYSTEM": "ប្រព័ន្ធម៉ាក",
    "A simpler mark and a more considered interface": "និមិត្តសញ្ញាសាមញ្ញ និង Interface ដែលគិតគូរល្អជាងមុន",
    "The visual language is now more coherent: navy, ivory, brushed gold and calm motion.": "ភាសារូបរាងឥឡូវមានភាពស៊ីសង្វាក់ជាងមុន: Navy, Ivory, មាសបែប Brushed និងចលនាស្ងប់ស្ងាត់។",
    "Simple gold A monogram": "អក្សរ A មាសដ៏សាមញ្ញ",
    "No temple, skyline or landscape symbol. Just a clear, scalable hospitality mark.": "មិនមានប្រាសាទ, Skyline ឬទេសភាពស្ថាបត្យកម្មទេ។ មានតែនិមិត្តសញ្ញាបដិសណ្ឋារកិច្ចដែលច្បាស់ និងពង្រីកបានល្អ។",
    "Interaction decisions": "ការសម្រេចចិត្តលើ Interaction",
    "Balanced guest shell": "ផ្ទាំងភ្ញៀវដែលមានតុល្យភាព",
    "Ivory body, stable navy glass navigation and gold hairline.": "ផ្ទៃ Ivory, Navigation ពណ៌ Navy មានស្ថេរភាព និងបន្ទាត់មាសស្តើង។",
    "Welcome modal": "Modal ស្វាគមន៍",
    "Blurred backdrop, centered offer card and one-session dismissal.": "ផ្ទៃខាងក្រោយព្រិល, Card ផ្តល់ជូននៅកណ្តាល និងបិទបានម្តងក្នុងមួយ Session។",
    "Motion with control": "ចលនាដែលមានការគ្រប់គ្រង",
    "Page transitions, press feedback and reduced-motion fallbacks.": "ការប្តូរទំព័រ, ប្រតិកម្មពេលចុច និងជម្រើសសម្រាប់អ្នកដែលមិនចង់បានចលនាច្រើន។",
    "Accessible behavior": "អាកប្បកិរិយាដែលងាយប្រើសម្រាប់គ្រប់គ្នា",
    "Escape close, focus restore/trap, live status and responsive controls.": "បិទដោយ Escape, រក្សា/ចាប់ Focus, ស្ថានភាព Live និង Controls ដែលឆ្លើយតបល្អ។",
    "DEPLOYMENT PATH": "ផ្លូវដាក់ឱ្យប្រើ",
    "SQLite fallback, MySQL target, Navicat visibility": "SQLite សម្រាប់ប្រើភ្លាមៗ, MySQL ជាគោលដៅ និងមើលទិន្នន័យតាម Navicat",
    "Windows setup scripts now carry the project from a local snapshot to a user-owned MySQL database.": "ស្គ្រីប Setup របស់ Windows អាចនាំគម្រោងពី Snapshot មូលដ្ឋាន ទៅកាន់ Database MySQL របស់ម្ចាស់ប្រព័ន្ធ។",
    "Current SQLite": "SQLite បច្ចុប្បន្ន",
    "Transfer helper": "ឧបករណ៍ជំនួយផ្ទេរ",
    "MySQL / MariaDB": "MySQL / MariaDB",
    "same connection": "ការភ្ជាប់ដូចគ្នា",
    "One-click Windows path": "ផ្លូវដំណើរការ Windows ដោយចុចម្តង",
    "install.bat  →  migrate  →  import if empty  →  audit\nstart.bat    →  connect  →  open admin preview": "install.bat  →  Migrate  →  Import បើទទេ  →  Audit\nstart.bat    →  ភ្ជាប់  →  បើក Preview សម្រាប់ Admin",
    "Automatic defaults": "តម្លៃកំណត់ស្វ័យប្រវត្តិ",
    "Use mysql.local.bat for a private password and Navicat-specific values.": "ប្រើ mysql.local.bat សម្រាប់ Password ឯកជន និងតម្លៃកំណត់ជាក់លាក់របស់ Navicat។",
    "QUALITY STATUS": "ស្ថានភាពគុណភាព",
    "Release checks are green": "ការត្រួតពិនិត្យកំណែនេះបានជោគជ័យ",
    "Automated validation and live route checks cover the latest design, booking and deployment changes.": "ការផ្ទៀងផ្ទាត់ស្វ័យប្រវត្តិ និងការត្រួតពិនិត្យ Route ផ្ទាល់ គ្របដណ្តប់ការកែប្រែ Design, ការកក់ និងការដាក់ឱ្យប្រើចុងក្រោយ។",
    "QA CLEAN": "QA ស្អាត",
    "Latest automated run": "លទ្ធផលស្វ័យប្រវត្តិចុងក្រោយ",
    "Django tests": "តេស្ត Django",
    "ledger invariants": "Invariant របស់ Ledger",
    "pending migrations": "Migration នៅសល់",
    "Hotels route": "Route សណ្ឋាគារ",
    "No 500 response or traceback observed during the latest live crawl.": "មិនបានរកឃើញ Response 500 ឬ Traceback ក្នុងការត្រួតពិនិត្យ Live ចុងក្រោយទេ។",
    "Verified coverage": "ផ្នែកដែលបានផ្ទៀងផ្ទាត់",
    "•  Anonymous guest shell, hotel collection, rooms, map, auth pages and designed 404\n•  Guest, reception, manager, housekeeping and admin route/role gates\n•  Room Board SSE stream with valid JSON for all 369 rooms\n•  Welcome offer eligibility, one-time consumption and financial snapshot regression\n•  JavaScript syntax, local CSS assets, logo assets and versioned cache paths": "•  ផ្ទាំង Guest មិនទាន់ Login, Collection, បន្ទប់, ផែនទី, ទំព័រ Auth និង 404 ដែលបានរចនា\n•  Route/Role gate សម្រាប់ Guest, Reception, Manager, Housekeeping និង Admin\n•  Room Board SSE Stream ដែលបញ្ជូន JSON ត្រឹមត្រូវសម្រាប់បន្ទប់ ៣៦៩\n•  សិទ្ធិទទួល Offer, ការប្រើបានម្តង និង Regression របស់ Financial Snapshot\n•  Syntax JavaScript, CSS Assets មូលដ្ឋាន, Logo Assets និង Cache Path តាម Version",
    "Ready for the next user-owned MySQL/Navicat connection step.": "ត្រៀមរួចសម្រាប់ជំហានបន្ទាប់: ភ្ជាប់ MySQL/Navicat របស់ម្ចាស់ប្រព័ន្ធ។",
    "RELEASE HANDOFF": "ការប្រគល់កំណែ",
    "Current version: ready for configuration": "កំណែបច្ចុប្បន្ន: ត្រៀមសម្រាប់កំណត់រចនាសម្ព័ន្ធ",
    "The product build is complete; the remaining environment step is connecting the owner’s MySQL/MariaDB server.": "ការបង្កើតផលិតផលបានបញ្ចប់ហើយ; ជំហានបរិស្ថានដែលនៅសល់ គឺភ្ជាប់ទៅ Server MySQL/MariaDB របស់ម្ចាស់ប្រព័ន្ធ។",
    "Recommended handoff": "ជំហានប្រគល់ដែលណែនាំ",
    "Start MySQL/MariaDB": "ចាប់ផ្តើម MySQL/MariaDB",
    "Use the same server details configured in Navicat.": "ប្រើព័ត៌មាន Server ដូចគ្នាដែលបានកំណត់ក្នុង Navicat។",
    "Copy mysql.local.bat.example": "ចម្លង mysql.local.bat.example",
    "Enter the private password and any custom host/port.": "បញ្ចូល Password ឯកជន និង Host/Port ផ្ទាល់ខ្លួន ប្រសិនបើមាន។",
    "Run install.bat once": "ដំណើរការ install.bat ម្តង",
    "Migrate schema, preserve the SQLite snapshot and audit.": "ផ្ទេរ Schema, រក្សា SQLite Snapshot និងដំណើរការ Audit។",
    "Run start.bat": "ដំណើរការ start.bat",
    "Launch the app and open the admin preview automatically.": "បើកកម្មវិធី និងបើក Admin Preview ដោយស្វ័យប្រវត្តិ។",
    "10 sanctuaries · Cambodia": "សណ្ឋាគារ ១០ កន្លែង · កម្ពុជា",
    "Latest build\n27 September 2026": "Build ចុងក្រោយ\n២៧ កញ្ញា ២០២៦",
    # Operations guide slides
    "PROJECT JOURNEY": "ដំណើររបស់គម្រោង",
    "From launch to a settled stay": "ពីពេលចាប់ផ្តើម រហូតដល់បិទការស្នាក់នៅ",
    "The system connects the public guest experience to daily hotel operations and final financial reporting.": "ប្រព័ន្ធភ្ជាប់បទពិសោធន៍ភ្ញៀវសាធារណៈ ជាមួយប្រតិបត្តិការប្រចាំថ្ងៃ និងរបាយការណ៍ហិរញ្ញវត្ថុចុងក្រោយ។",
    "Launch": "ចាប់ផ្តើម",
    "SQLite starts automatically. Public homepage opens signed out.": "SQLite ចាប់ផ្តើមដោយស្វ័យប្រវត្តិ។ Homepage សាធារណៈបើកដោយមិន Login។",
    "Discover": "ស្វែងយល់",
    "Guest explores properties, rooms, photos and maps.": "ភ្ញៀវស្វែងរកសណ្ឋាគារ, បន្ទប់, រូបថត និងផែនទី។",
    "Price": "គណនាតម្លៃ",
    "Availability, tiered rates, offer, service and tax are calculated.": "ប្រព័ន្ធគណនា Availability, តម្លៃតាមកម្រិត, Offer, សេវា និងពន្ធ។",
    "Availability, tiered rates, offer and 5% tax are calculated.": "ប្រព័ន្ធគណនា Availability, តម្លៃតាមកម្រិត, Offer និងពន្ធ ៥%។",
    "Booking, folio, payment and rate snapshot are created.": "ប្រព័ន្ធបង្កើតការកក់, Folio, ការទូទាត់ និង Rate Snapshot។",
    "Operate": "ប្រតិបត្តិការ",
    "Front desk assigns rooms; housekeeping and maintenance work.": "Front desk ចាត់បន្ទប់; Housekeeping និង Maintenance ចាប់ផ្តើមការងារ។",
    "Stay": "ស្នាក់នៅ",
    "Guest is checked in, in-house, then checked out.": "ភ្ញៀវ Check-in, ស្នាក់នៅក្នុងសណ្ឋាគារ ហើយបន្ទាប់មក Check-out។",
    "Room becomes dirty, task is created and folio is settled.": "បន្ទប់ត្រូវបានសម្គាល់ថាត្រូវសម្អាត, កិច្ចការត្រូវបានបង្កើត និង Folio ត្រូវបានបិទបញ្ជី។",
    "Learn": "វិភាគ",
    "Managers review revenue, occupancy, channel mix and audit data.": "Manager ពិនិត្យចំណូល, Occupancy, ប្រភពការកក់ និងទិន្នន័យ Audit។",
    "One source of truth: the booking snapshot travels with the guest from reservation to folio to reporting.": "ប្រភពទិន្នន័យតែមួយ: Booking Snapshot តាមដំណើរភ្ញៀវពីការកក់ ទៅ Folio និងរបាយការណ៍។",
    "BOOKING CHANNELS": "បណ្តាញចូលនៃការកក់",
    "Four ways a reservation can enter the system": "វិធី ៤ យ៉ាងដែលការកក់អាចចូលក្នុងប្រព័ន្ធ",
    "The PMS uses one operational workflow while preserving where each booking came from.": "PMS ប្រើលំហូរប្រតិបត្តិការតែមួយ ខណៈពេលរក្សាទុកប្រភពដើមនៃការកក់នីមួយៗ។",
    "WEBSITE": "គេហទំព័រ",
    "Website": "គេហទំព័រ",
    "Source code: web  ·  Owner: Guest self-service": "Source code: web  ·  អ្នកទទួលខុសត្រូវ: ភ្ញៀវកក់ដោយខ្លួនឯង",
    "•  Guest searches and books\n•  Eligible registered guest receives 10% first-stay offer\n•  Online guest can view My Stays and cancel": "•  ភ្ញៀវស្វែងរក និងកក់ដោយខ្លួនឯង\n•  ភ្ញៀវដែលចុះឈ្មោះ និងមានសិទ្ធិ ទទួលបាន Offer ១០% សម្រាប់ការស្នាក់នៅលើកដំបូង\n•  ភ្ញៀវ Online អាចមើល My Stays និងលុបចោល",
    "WALK-IN": "មកផ្ទាល់",
    "Walk-in": "មកផ្ទាល់ (Walk-in)",
    "Source code: walk_in  ·  Owner: Front desk": "Source code: walk_in  ·  អ្នកទទួលខុសត្រូវ: Front desk",
    "•  No online account required\n•  Receptionist enters guest and stay details\n•  Cash, card, wallet or bank transfer can be recorded": "•  មិនចាំបាច់មានគណនី Online\n•  Receptionist បញ្ចូលព័ត៌មានភ្ញៀវ និងការស្នាក់នៅ\n•  អាចកត់ត្រា Cash, Card, Wallet ឬ Bank transfer",
    "PHONE": "ទូរស័ព្ទ",
    "Phone": "ទូរស័ព្ទ",
    "Source code: phone  ·  Owner: Front desk": "Source code: phone  ·  អ្នកទទួលខុសត្រូវ: Front desk",
    "•  Receptionist books for the caller\n•  Availability and pricing are validated live\n•  Guest contact details are stored on the reservation": "•  Receptionist កក់ជំនួសអ្នកហៅទូរស័ព្ទ\n•  Availability និងតម្លៃត្រូវបានផ្ទៀងផ្ទាត់ភ្លាមៗ\n•  ព័ត៌មានទំនាក់ទំនងភ្ញៀវត្រូវបានរក្សាទុកក្នុងការកក់",
    "OTA": "OTA",
    "OTA / Agency": "OTA / ភ្នាក់ងារ",
    "Source code: ota  ·  Owner: Manual staff entry today": "Source code: ota  ·  អ្នកទទួលខុសត្រូវ: បុគ្គលិកបញ្ចូលដោយដៃបច្ចុប្បន្ន",
    "•  Source is preserved for reporting\n•  External reference can be kept in notes\n•  Automatic channel-manager sync is a future step": "•  ប្រភពត្រូវបានរក្សាទុកសម្រាប់របាយការណ៍\n•  អាចរក្សាលេខយោងខាងក្រៅក្នុង Notes\n•  ការភ្ជាប់ស្វ័យប្រវត្តិជាមួយ Channel Manager ជាជំហានអនាគត",
    "Current boundary: OTA/agency reservations are recorded manually; live Booking.com/Expedia-style synchronization is not yet implemented.": "កម្រិតបច្ចុប្បន្ន: ការកក់ OTA/Agency ត្រូវបានបញ្ចូលដោយដៃ; ការធ្វើសមកាលកម្មផ្ទាល់បែប Booking.com/Expedia មិនទាន់បានអនុវត្តទេ។",
    "ROLES & ACCESS": "តួនាទី និងសិទ្ធិចូលប្រើ",
    "Who can do what": "អ្នកណាអាចធ្វើអ្វីបានខ្លះ",
    "The release has guest access, staff access, and manager/administrator business gates.": "កំណែនេះមានសិទ្ធិសម្រាប់ Guest, Staff និងច្រកសិទ្ធិអាជីវកម្មសម្រាប់ Manager/Administrator។",
    "ROLE": "តួនាទី",
    "PRIMARY JOB": "ការងារសំខាន់",
    "CAN ACCESS": "អ្វីដែលអាចចូលប្រើ",
    "BOUNDARY / NOTES": "ព្រំដែន / ចំណាំ",
    "Guest": "Guest / ភ្ញៀវ",
    "Book and manage own stay": "កក់ និងគ្រប់គ្រងការស្នាក់នៅរបស់ខ្លួន",
    "Public site · My Stays · self-service cancellation": "គេហទំព័រសាធារណៈ · My Stays · លុបចោលដោយខ្លួនឯង",
    "Cannot enter staff PMS or view other guests.": "មិនអាចចូល PMS របស់បុគ្គលិក ឬមើលភ្ញៀវផ្សេងទៀតបានទេ។",
    "Receptionist": "Receptionist / ផ្នែកទទួលភ្ញៀវ",
    "Front desk and guest movement": "Front desk និងចលនារបស់ភ្ញៀវ",
    "Reservations · calendar · channels · check-in/out · payments": "ការកក់ · ប្រតិទិន · បណ្តាញ · Check-in/out · ការទូទាត់",
    "Operational staff access; no rates, reports or analytics gate.": "មានសិទ្ធិបុគ្គលិកប្រតិបត្តិការ; មិនមានសិទ្ធិគ្រប់គ្រង Rate, Report ឬ Analytics ទេ។",
    "Housekeeping": "Housekeeping / ផ្នែកសម្អាត",
    "Room readiness and tasks": "រៀបចំបន្ទប់ឱ្យរួចរាល់ និងគ្រប់គ្រងកិច្ចការ",
    "Room Board · housekeeping · maintenance": "Room Board · Housekeeping · Maintenance",
    "Primary job is rooms; current broad staff gate also permits some shared staff pages.": "ការងារសំខាន់គឺបន្ទប់; តាមបច្ចុប្បន្ន ច្រក Staff ទូលំទូលាយក៏អនុញ្ញាតឱ្យចូលទំព័ររួមមួយចំនួនដែរ។",
    "Manager": "Manager / អ្នកគ្រប់គ្រង",
    "Commercial and property performance": "លទ្ធផលពាណិជ្ជកម្ម និងប្រតិបត្តិការសណ្ឋាគារ",
    "All staff areas + rates · finance reports · analytics": "គ្រប់ផ្នែក Staff + Rate · របាយការណ៍ហិរញ្ញវត្ថុ · Analytics",
    "Manager-level gate required for rate rules and business reporting.": "ត្រូវការសិទ្ធិកម្រិត Manager សម្រាប់ច្បាប់ Rate និងរបាយការណ៍អាជីវកម្ម។",
    "Administrator": "Administrator / អ្នកគ្រប់គ្រងប្រព័ន្ធ",
    "System and hotel administration": "គ្រប់គ្រងប្រព័ន្ធ និងសណ្ឋាគារ",
    "Everything + Django admin · users · configuration": "គ្រប់យ៉ាង + Django Admin · អ្នកប្រើ · ការកំណត់រចនាសម្ព័ន្ធ",
    "Highest application role; manager privileges included.": "ជាតួនាទីខ្ពស់បំផុតក្នុងកម្មវិធី; រួមមានសិទ្ធិរបស់ Manager។",
    "Implementation note: the current RBAC has two major technical gates — staff and manager/admin. Fine-grained housekeeping-only restrictions can be added later if required.": "ចំណាំអំពីការអនុវត្ត: RBAC បច្ចុប្បន្នមានច្រកបច្ចេកទេសធំ ២ — Staff និង Manager/Admin។ ប្រសិនបើត្រូវការ អាចបន្ថែមការកំណត់សិទ្ធិឱ្យ Housekeeping ដាច់ដោយឡែកបាននៅពេលក្រោយ។",
    "SERVICE MAP": "ផែនទីសេវាកម្ម",
    "The right person acts at the right moment": "មនុស្សត្រឹមត្រូវ ធ្វើសកម្មភាពនៅពេលត្រឹមត្រូវ",
    "This is the operational handoff from guest request to hotel team to financial close.": "នេះជាលំហូរប្រតិបត្តិការពីសំណើភ្ញៀវ ទៅក្រុមសណ្ឋាគារ និងបិទបញ្ជីហិរញ្ញវត្ថុ។",
    "Discover and request": "ស្វែងរក និងស្នើសុំ",
    "Website, property picker, room search, welcome offer, payment intent and My Stays.": "គេហទំព័រ, ជ្រើសសណ្ឋាគារ, ស្វែងរកបន្ទប់, Welcome Offer, បំណងទូទាត់ និង My Stays។",
    "Reception": "Reception / ផ្នែកទទួលភ្ញៀវ",
    "Convert and control": "បម្លែងការស្នើសុំ និងគ្រប់គ្រង",
    "Walk-in, phone and OTA/agency reservations; availability; room assignment; check-in/out.": "ការកក់ Walk-in, ទូរស័ព្ទ និង OTA/Agency; Availability; ចាត់បន្ទប់; Check-in/out។",
    "Prepare and restore": "រៀបចំ និងស្តារស្ថានភាពបន្ទប់",
    "Vacant-clean status, task start/complete, post-checkout room recovery and maintenance reporting.": "ស្ថានភាព Vacant-clean, ចាប់ផ្តើម/បញ្ចប់កិច្ចការ, រៀបចំបន្ទប់ក្រោយ Check-out និងរាយការណ៍ Maintenance។",
    "Finance": "ហិរញ្ញវត្ថុ",
    "Record and settle": "កត់ត្រា និងបិទបញ្ជី",
    "Folio, invoice, full-stay payment, refund entry and outstanding balance.": "Folio, Invoice, ការទូទាត់ពេញការស្នាក់នៅ, កំណត់ត្រា Refund និងសមតុល្យនៅសល់។",
    "Review and improve": "ពិនិត្យ និងកែលម្អ",
    "Rates, rules, occupancy, ADR, RevPAR, revenue, source mix and property performance.": "Rate, ច្បាប់, Occupancy, ADR, RevPAR, ចំណូល, ប្រភពការកក់ និងលទ្ធផលសណ្ឋាគារ។",
    "Guest-facing hospitality services shown on the public site include airport transfers, 24-hour concierge and limousine service.": "សេវាបដិសណ្ឋារកិច្ចដែលបង្ហាញលើគេហទំព័ររួមមាន Airport Transfer, Concierge ២៤ ម៉ោង និងសេវា Limousine។",
    "HANDOFF CLARITY": "ភាពច្បាស់លាស់ពេលប្រគល់",
    "What is live now — and what comes next": "អ្វីដែលដំណើរការហើយ — និងអ្វីដែលត្រូវធ្វើបន្ទាប់",
    "The current release is ready for use; the next work is mostly environment integration and deeper hotel-system maturity.": "កំណែបច្ចុប្បន្នត្រៀមប្រើប្រាស់ហើយ; ការងារបន្ទាប់ផ្តោតលើការភ្ជាប់បរិស្ថាន និងការធ្វើឱ្យប្រព័ន្ធសណ្ឋាគារកាន់តែពេញលេញ។",
    "LIVE IN THIS RELEASE": "មាននៅក្នុងកំណែនេះ",
    "•  Public guest website and responsive premium interface\n•  10 properties · 24 room types · 369 rooms\n•  Tiered $30–$460 rates, 5% tax only and real one-time 10% offer\n•  Website, walk-in, phone and OTA/agency source handling\n•  Reservations, room board, housekeeping, maintenance and check-in/out\n•  Folios, payments, finance and analytics with snapshot integrity\n•  SQLite automatic launch; optional MySQL/MariaDB and Navicat kit\n•  23/23 tests and 2,160 financial invariants passed": "•  គេហទំព័រសាធារណៈសម្រាប់ភ្ញៀវ និង Interface ប្រណិតដែលឆ្លើយតបល្អ\n•  សណ្ឋាគារ ១០ · បន្ទប់ ២៤ ប្រភេទ · បន្ទប់ ៣៦៩\n•  តម្លៃតាមកម្រិត $30–$460, ពន្ធតែ ៥% និង Offer ១០% ប្រើបានម្តងពិតប្រាកដ\n•  គាំទ្រប្រភព Website, Walk-in, Phone និង OTA/Agency\n•  ការកក់, Room Board, Housekeeping, Maintenance និង Check-in/out\n•  Folio, ការទូទាត់, ហិរញ្ញវត្ថុ និង Analytics ដែលរក្សា Snapshot ត្រឹមត្រូវ\n•  SQLite ចាប់ផ្តើមស្វ័យប្រវត្តិ; មានជម្រើស MySQL/MariaDB និង Navicat Kit\n•  តេស្ត ២៣/២៣ និង Financial Invariant ២,១៦០ បានជោគជ័យ",
    "RECOMMENDED NEXT": "អ្វីដែលណែនាំឱ្យធ្វើបន្ទាប់",
    "•  Configure and verify the user-owned MySQL/Navicat connection\n•  Add automatic OTA/channel-manager synchronization\n•  Tighten housekeeping/reception permissions if needed\n•  Add folio line items for minibar, spa, restaurant and transport\n•  Add rate plans, group blocks, waitlist and room moves\n•  Add English/KH localization and broader browser testing": "•  កំណត់ និងផ្ទៀងផ្ទាត់ការភ្ជាប់ MySQL/Navicat របស់ម្ចាស់ប្រព័ន្ធ\n•  បន្ថែមការធ្វើសមកាលកម្មស្វ័យប្រវត្តិជាមួយ OTA/Channel Manager\n•  បែងចែកសិទ្ធិ Housekeeping/Reception ឱ្យតឹងជាងមុន ប្រសិនបើត្រូវការ\n•  បន្ថែម Folio Line Item សម្រាប់ Minibar, Spa, ភោជនីយដ្ឋាន និងដឹកជញ្ជូន\n•  បន្ថែម Rate Plan, Group Block, Waitlist និងផ្លាស់ប្តូរបន្ទប់\n•  បន្ថែមភាសាខ្មែរ/អង់គ្លេស និងការធ្វើតេស្ត Browser ឱ្យទូលំទូលាយ",
    "Final environment step: owner-controlled database verification.": "ជំហានបរិស្ថានចុងក្រោយ: ផ្ទៀងផ្ទាត់ Database ដែលម្ចាស់ប្រព័ន្ធគ្រប់គ្រង។",
    "QUICK REFERENCE": "ឯកសារយោងរហ័ស",
    "Aurelia Collection in one sentence": "Aurelia Collection ក្នុងមួយប្រយោគ",
    "A premium guest booking experience connected to a practical, auditable hotel operating system.": "បទពិសោធន៍កក់សម្រាប់ភ្ញៀវប្រកបដោយភាពប្រណិត ដែលភ្ជាប់ជាមួយប្រព័ន្ធប្រតិបត្តិការសណ្ឋាគារដែលអាចប្រើការ និង Audit បាន។",
    "Guest website  →  reservation channel  →  availability & pricing  →  room operations  →  folio & payment  →  management insight": "គេហទំព័រភ្ញៀវ  →  បណ្តាញកក់  →  Availability និងតម្លៃ  →  ប្រតិបត្តិការបន្ទប់  →  Folio និងការទូទាត់  →  ព័ត៌មានសម្រាប់អ្នកគ្រប់គ្រង",
    "Guests": "ភ្ញៀវ",
    "Discover, book, pay, view My Stays and cancel.": "ស្វែងរក, កក់, ទូទាត់, មើល My Stays និងលុបចោល។",
    "Front desk": "Front desk",
    "Control reservations, room movement and guest arrival.": "គ្រប់គ្រងការកក់, ចលនាបន្ទប់ និងការមកដល់របស់ភ្ញៀវ។",
    "Operations": "ប្រតិបត្តិការ",
    "Keep rooms clean, sellable, maintained and visible.": "រក្សាបន្ទប់ឱ្យស្អាត, អាចលក់បាន, មានការថែទាំ និងមើលឃើញស្ថានភាព។",
    "Management": "ការគ្រប់គ្រង",
    "Control rates, finance, analytics and system setup.": "គ្រប់គ្រង Rate, ហិរញ្ញវត្ថុ, Analytics និងការកំណត់ប្រព័ន្ធ។",
    "Next action: configure the user-owned MySQL/Navicat connection and verify the live environment.": "សកម្មភាពបន្ទាប់: កំណត់ការភ្ជាប់ MySQL/Navicat របស់ម្ចាស់ប្រព័ន្ធ និងផ្ទៀងផ្ទាត់បរិស្ថាន Live។",
    "PROJECT PURPOSE": "គោលបំណងគម្រោង",
    "Why the hotel management project exists": "ហេតុអ្វីបានជាមានគម្រោងគ្រប់គ្រងសណ្ឋាគារ",
    "One system connects guest demand, hotel operations and financial control across a multi-property collection.": "ប្រព័ន្ធតែមួយភ្ជាប់តម្រូវការភ្ញៀវ, ប្រតិបត្តិការសណ្ឋាគារ និងការគ្រប់គ្រងហិរញ្ញវត្ថុ សម្រាប់សណ្ឋាគារជាច្រើន។",
    "The problem solved": "បញ្ហាដែលគម្រោងនេះដោះស្រាយ",
    "A hotel should not need separate answers for booking, rooms, housekeeping and money.": "សណ្ឋាគារមិនគួរត្រូវប្រើប្រព័ន្ធខុសៗគ្នា សម្រាប់ការកក់, បន្ទប់, Housekeeping និងហិរញ្ញវត្ថុឡើយ។",
    "•  One booking record from any channel\n•  One availability and pricing engine\n•  One room-status and operations view\n•  One folio and financial history": "•  កំណត់ត្រាកក់តែមួយ ពីគ្រប់បណ្តាញ\n•  ម៉ាស៊ីន Availability និង Pricing តែមួយ\n•  ទិដ្ឋភាពស្ថានភាពបន្ទប់ និងប្រតិបត្តិការតែមួយ\n•  Folio និងប្រវត្តិហិរញ្ញវត្ថុតែមួយ",
    "Project scope": "វិសាលភាពគម្រោង",
    "Guest": "ភ្ញៀវ",
    "Public collection, rooms, booking, payment intent, My Stays": "Collection សាធារណៈ, បន្ទប់, ការកក់, បំណងទូទាត់ និង My Stays",
    "Front office": "Front office",
    "Reservations, calendar, check-in, checkout, room assignment": "ការកក់, ប្រតិទិន, Check-in, Check-out និងចាត់បន្ទប់",
    "Operations": "ប្រតិបត្តិការ",
    "Room Board, housekeeping, maintenance and live status": "Room Board, Housekeeping, Maintenance និងស្ថានភាពផ្ទាល់",
    "Business": "អាជីវកម្ម",
    "Rates, folios, payments, reports, analytics and roles": "Rate, Folio, ការទូទាត់, របាយការណ៍, Analytics និងតួនាទី",
    "Deployment": "ការដាក់ឱ្យប្រើ",
    "SQLite default, MySQL/MariaDB option, Navicat export": "SQLite ជាលំនាំដើម, ជម្រើស MySQL/MariaDB និងការនាំចេញ Navicat",
    "SYSTEM MODEL": "គំរូប្រព័ន្ធ",
    "The hotel data that powers the PMS": "ទិន្នន័យសណ្ឋាគារដែលជំរុញ PMS",
    "Every operating screen is built around properties, rooms, bookings, folios and people.": "គ្រប់ផ្ទាំងប្រតិបត្តិការ ស្ថាបនាឡើងជុំវិញសណ្ឋាគារ, បន្ទប់, ការកក់, Folio និងអ្នកប្រើប្រាស់។",
    "Properties": "សណ្ឋាគារ / អចលនទ្រព្យ",
    "Hotels and destinations in the collection": "សណ្ឋាគារ និងគោលដៅក្នុង Collection",
    "Sellable categories with base rates": "ប្រភេទបន្ទប់ដែលអាចលក់បាន ជាមួយតម្លៃមូលដ្ឋាន",
    "Rooms": "បន្ទប់",
    "Physical inventory assigned to room types": "ស្តុកបន្ទប់ជាក់ស្តែង ដែលភ្ជាប់ជាមួយប្រភេទបន្ទប់",
    "Reservations across all channels": "ការកក់ពីគ្រប់បណ្តាញ",
    "One financial record for each stay": "កំណត់ត្រាហិរញ្ញវត្ថុមួយសម្រាប់ការស្នាក់នៅនីមួយៗ",
    "Charges and refunds with status": "ការគិតប្រាក់ និង Refund ជាមួយស្ថានភាព",
    "Core relationship": "ទំនាក់ទំនងស្នូល",
    "Property  →  Room type  →  Room inventory  →  Booking  →  Folio  →  Payment": "សណ្ឋាគារ  →  ប្រភេទបន្ទប់  →  ស្តុកបន្ទប់  →  ការកក់  →  Folio  →  ការទូទាត់",
    "Users and roles control who can operate each part of this chain.": "Users និង Roles កំណត់ថា អ្នកណាអាចប្រតិបត្តិការផ្នែកនីមួយៗក្នុងខ្សែសង្វាក់នេះ។",
    "GUEST FLOW": "លំហូររបស់ភ្ញៀវ",
    "Website booking in six simple steps": "ការកក់តាមគេហទំព័រ ក្នុង ៦ ជំហានសាមញ្ញ",
    "The guest sees a clear path; the system performs the complex validation behind it.": "ភ្ញៀវឃើញផ្លូវដែលច្បាស់; ប្រព័ន្ធធ្វើការផ្ទៀងផ្ទាត់ស្មុគស្មាញនៅពីក្រោយ។",
    "View properties, maps, photos and hotel atmosphere.": "មើលសណ្ឋាគារ, ផែនទី, រូបថត និងបរិយាកាសសណ្ឋាគារ។",
    "Search": "ស្វែងរក",
    "Choose dates, guests and rooms; availability is checked.": "ជ្រើសកាលបរិច្ឆេទ, ចំនួនភ្ញៀវ និងបន្ទប់; ប្រព័ន្ធពិនិត្យ Availability។",
    "Select": "ជ្រើសរើស",
    "Compare room types, amenities and nightly rates.": "ប្រៀបធៀបប្រភេទបន្ទប់, សេវាបន្ថែម និងតម្លៃក្នុងមួយយប់។",
    "Offer": "ការផ្តល់ជូន",
    "Eligible registered guests unlock a one-time 10% discount.": "ភ្ញៀវដែលបានចុះឈ្មោះ និងមានសិទ្ធិ អាចបើកការបញ្ចុះតម្លៃ ១០% ដែលប្រើបានម្តង។",
    "Book": "កក់",
    "Enter contact details, requests and payment method.": "បញ្ចូលព័ត៌មានទំនាក់ទំនង, សំណើពិសេស និងវិធីទូទាត់។",
    "Manage": "គ្រប់គ្រង",
    "Receive confirmation, view My Stays and cancel if allowed.": "ទទួលការបញ្ជាក់, មើល My Stays និងលុបចោល ប្រសិនបើគោលការណ៍អនុញ្ញាត។",
    "Result: the guest gets a confirmation while the PMS receives a validated reservation, folio and financial snapshot.": "លទ្ធផល: ភ្ញៀវទទួលបានការបញ្ជាក់ ខណៈ PMS ទទួលបានការកក់ដែលបានផ្ទៀងផ្ទាត់, Folio និង Financial Snapshot។",
    "DAILY OPERATIONS": "ប្រតិបត្តិការប្រចាំថ្ងៃ",
    "A simple playbook for the hotel team": "សៀវភៅណែនាំសាមញ្ញសម្រាប់ក្រុមសណ្ឋាគារ",
    "The PMS follows the property’s day: arrivals, rooms, guest service, departures and close.": "PMS ដើរតាមវដ្តប្រចាំថ្ងៃរបស់សណ្ឋាគារ: ភ្ញៀវមកដល់, បន្ទប់, សេវាភ្ញៀវ, ភ្ញៀវចាកចេញ និងបិទថ្ងៃ។",
    "Morning": "ពេលព្រឹក",
    "Prepare": "រៀបចំ",
    "Review arrivals, departures, occupancy, dirty rooms and open tasks.": "ពិនិត្យភ្ញៀវមកដល់, ភ្ញៀវចាកចេញ, Occupancy, បន្ទប់ត្រូវសម្អាត និងកិច្ចការដែលនៅបើក។",
    "Arrival": "ពេលមកដល់",
    "Check in": "Check-in",
    "Verify booking, assign vacant-clean room and move guest to in-house.": "ផ្ទៀងផ្ទាត់ការកក់, ចាត់បន្ទប់ Vacant-clean និងប្តូរស្ថានភាពភ្ញៀវទៅ In-house។",
    "During stay": "អំឡុងពេលស្នាក់នៅ",
    "Serve": "បម្រើ",
    "Monitor room status, requests, maintenance and folio activity.": "តាមដានស្ថានភាពបន្ទប់, សំណើភ្ញៀវ, Maintenance និងសកម្មភាព Folio។",
    "Departure": "ពេលចាកចេញ",
    "Check out": "Check-out",
    "Settle folio, release room and automatically create housekeeping work.": "បិទ Folio, ដោះលែងបន្ទប់ និងបង្កើតកិច្ចការ Housekeeping ដោយស្វ័យប្រវត្តិ។",
    "Night close": "បិទប្រតិបត្តិការពេលយប់",
    "Review": "ពិនិត្យឡើងវិញ",
    "Run nightly operations, flag no-shows and produce occupancy snapshot.": "ដំណើរការ Nightly Operations, សម្គាល់ No-show និងបង្កើត Occupancy Snapshot។",
    "Every action updates the same reservation, room status or folio — no disconnected spreadsheets required.": "សកម្មភាពគ្រប់យ៉ាងធ្វើបច្ចុប្បន្នភាពលើការកក់, ស្ថានភាពបន្ទប់ ឬ Folio ដដែល — មិនចាំបាច់ប្រើ Spreadsheet ដាច់ៗគ្នាទេ។",
    "CONTROLS": "ការគ្រប់គ្រង និងការពារ",
    "How the project protects the numbers": "របៀបដែលគម្រោងការពារភាពត្រឹមត្រូវនៃលេខ",
    "The price shown, the folio issued and the historical total are designed to agree.": "តម្លៃដែលបង្ហាញ, Folio ដែលចេញ និងចំនួនសរុបប្រវត្តិសាស្ត្រ ត្រូវបានរចនាឱ្យស្របគ្នា។",
    "Pricing engine": "ម៉ាស៊ីនគណនាតម្លៃ",
    "Base rate + active rate rule + weekend rule, multiplied by nights and rooms.": "តម្លៃមូលដ្ឋាន + ច្បាប់ Rate ដែលសកម្ម + ច្បាប់ Weekend បន្ទាប់មកគុណនឹងចំនួនយប់ និងបន្ទប់។",
    "Charges": "ការគិតប្រាក់",
    "5% service charge and 10% tax are calculated after the discounted room subtotal.": "សេវា ៥% និងពន្ធ ១០% ត្រូវបានគណនាបន្ទាប់ពី Subtotal បន្ទប់ដែលបានបញ្ចុះតម្លៃ។",
    "The Cambodia demo uses 5% tax only; there is no separate service charge.": "Demo កម្ពុជាប្រើពន្ធតែ ៥% ប៉ុណ្ណោះ; មិនមានសេវាគិតដាច់ដោយឡែកទេ។",
    "Snapshot": "Snapshot",
    "Rates, discount, service, tax and total are saved on the booking and never recomputed for history.": "Rate, បញ្ចុះតម្លៃ, សេវា, ពន្ធ និងសរុប ត្រូវបានរក្សាទុកក្នុងការកក់ ហើយមិនគណនាឡើងវិញសម្រាប់ប្រវត្តិទេ។",
    "Lifecycle safeguards": "ការការពារតាមវដ្តការកក់",
    "•  Availability is checked across every night and protected against double booking.\n•  Cancellation rules can retain the first night and create a refund ledger entry.\n•  High or critical maintenance requests block a room from being sold.": "•  Availability ត្រូវបានពិនិត្យគ្រប់យប់ និងការពារការកក់ស្ទួន។\n•  ច្បាប់លុបចោលអាចរក្សាទុកថ្លៃយប់ដំបូង និងបង្កើតកំណត់ត្រា Refund ក្នុង Ledger។\n•  សំណើ Maintenance កម្រិត High ឬ Critical នឹងរារាំងបន្ទប់មិនឱ្យលក់។",
    "QA proof": "ភស្តុតាង QA",
    "23/23 tests\n2,160 ledger invariants\n0 pending migrations": "តេស្ត ២៣/២៣\nLedger Invariant ២,១៦០\nMigration នៅសល់ ០",
    "FINAL SUMMARY": "សេចក្តីសង្ខេបចុងក្រោយ",
    "What the project delivers": "អ្វីដែលគម្រោងនេះផ្តល់ជូន",
    "A clear hotel operating model ready for owner-controlled database configuration and future channel integration.": "គំរូប្រតិបត្តិការសណ្ឋាគារដែលច្បាស់លាស់ ត្រៀមសម្រាប់ Database របស់ម្ចាស់ប្រព័ន្ធ និងការភ្ជាប់ Channel នៅពេលអនាគត។",
    "The project is ready to demonstrate": "គម្រោងត្រៀមសម្រាប់ Demo",
    "From a guest’s first click to a manager’s revenue report, the important hotel actions are connected.": "ចាប់ពីការចុចលើកដំបូងរបស់ភ្ញៀវ រហូតដល់របាយការណ៍ចំណូលរបស់ Manager សកម្មភាពសំខាន់ៗក្នុងសណ្ឋាគារត្រូវបានភ្ជាប់ជាមួយគ្នា។",
    "•  Guest experience is public, responsive and premium.\n•  Staff workflow is organized by role and operating moment.\n•  Money is preserved as a historical contract.\n•  SQLite runs automatically; MySQL/Navicat is the deployment path.": "•  បទពិសោធន៍ភ្ញៀវជាសាធារណៈ, ឆ្លើយតបល្អ និងមានភាពប្រណិត។\n•  លំហូរការងារបុគ្គលិកត្រូវបានរៀបចំតាមតួនាទី និងពេលប្រតិបត្តិការ។\n•  ទិន្នន័យប្រាក់ត្រូវបានរក្សាជាកិច្ចសន្យាប្រវត្តិសាស្ត្រ។\n•  SQLite ដំណើរការស្វ័យប្រវត្តិ; MySQL/Navicat ជាផ្លូវដាក់ឱ្យប្រើ។",
    "Final recommended sequence": "លំដាប់ជំហានចុងក្រោយដែលណែនាំ",
    "Use the public site": "ប្រើគេហទំព័រសាធារណៈ",
    "Start signed out and test the guest journey.": "ចាប់ផ្តើមដោយមិន Login ហើយសាកល្បងដំណើររបស់ភ្ញៀវ។",
    "Test each role": "សាកល្បងគ្រប់តួនាទី",
    "Confirm reception, housekeeping, manager and admin tasks.": "ផ្ទៀងផ្ទាត់ការងាររបស់ Reception, Housekeeping, Manager និង Admin។",
    "Configure MySQL": "កំណត់ MySQL",
    "Use the Navicat-owned server and private local settings.": "ប្រើ Server ដែលបានកំណត់ក្នុង Navicat និង Settings មូលដ្ឋានឯកជន។",
    "Add OTA integration": "បន្ថែមការភ្ជាប់ OTA",
    "Connect channel-manager APIs when ready.": "ភ្ជាប់ API របស់ Channel Manager នៅពេលត្រៀមរួច។",
    "Next action: verify the owner’s MySQL/Navicat connection.": "សកម្មភាពបន្ទាប់: ផ្ទៀងផ្ទាត់ការភ្ជាប់ MySQL/Navicat របស់ម្ចាស់ប្រព័ន្ធ។",
}

# Footer variants used by the source deck.
T["Aurelia Collection · Latest release recap · 27 September 2026"] = "Aurelia Collection · សង្ខេបកំណែចុងក្រោយ · ២៧ កញ្ញា ២០២៦"
T["Aurelia Collection · Project operations guide · 27 September 2026"] = "Aurelia Collection · មគ្គុទ្ទេសក៍ប្រតិបត្តិការ · ២៧ កញ្ញា ២០២៦"
T["Aurelia Collection · Hotel management project · 27 September 2026"] = "Aurelia Collection · គម្រោងគ្រប់គ្រងសណ្ឋាគារ · ២៧ កញ្ញា ២០២៦"


def replace_shape_text(shape, value):
    tf = shape.text_frame
    old_run = None
    old_align = None
    if tf.paragraphs:
        old_align = tf.paragraphs[0].alignment
        if tf.paragraphs[0].runs:
            old_run = tf.paragraphs[0].runs[0]
    size = old_run.font.size.pt if old_run and old_run.font.size else 11
    bold = bool(old_run.font.bold) if old_run else False
    italic = bool(old_run.font.italic) if old_run else False
    color = None
    if old_run:
        try:
            color = old_run.font.color.rgb
        except Exception:
            color = None
    # Khmer text is naturally a little wider; keep small body copy readable.
    n = len(value)
    if n > 190:
        size = min(size, 8.6)
    elif n > 120:
        size = min(size, 9.2)
    elif n > 75:
        size = min(size, 10.0)
    elif n > 45 and size >= 13:
        size = min(size, 12.5)
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if old_align is not None:
        p.alignment = old_align
    run = p.add_run()
    run.text = value
    run.font.name = "Noto Sans Khmer"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


untranslated = []
for slide in prs.slides:
    for shape in slide.shapes:
        if not hasattr(shape, "text_frame") or not shape.text.strip():
            continue
        original = shape.text
        translated = T.get(original)
        if translated is None:
            # Keep brand, numbers, technical filenames and arrows. Report any
            # other remaining Latin text for QA below.
            translated = original
            if any(c.isalpha() for c in original) and not original.startswith(("AURELIA", "aurelia", "mysql", "install", "start", "migrate")):
                untranslated.append(original)
        replace_shape_text(shape, translated)

prs.core_properties.title = "Aurelia Collection — គម្រោងគ្រប់គ្រងសណ្ឋាគារ"
prs.core_properties.subject = "សេចក្តីពិពណ៌នាពេញលេញជាភាសាខ្មែរ អំពីគម្រោងសណ្ឋាគារ, តួនាទី, បណ្តាញកក់ និងប្រតិបត្តិការ"
prs.core_properties.author = "Aurelia Collection"
prs.core_properties.keywords = "Khmer, hotel management, PMS, roles, website, walk-in, OTA, agency"
prs.core_properties.comments = "Natural Khmer translation of the validated 22-slide hotel management project deck."
prs.save(OUT)

print(OUT)
if untranslated:
    print("UNTRANSLATED_TEXT:")
    for value in sorted(set(untranslated)):
        print(value)
