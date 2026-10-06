from __future__ import annotations

from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
BASE = ROOT / "Aurelia_Collection_Project_Operations_Guide.pptx"
OUT = ROOT / "Aurelia_Collection_Hotel_Management_Project_Full_Deck.pptx"

NAVY = RGBColor(11, 20, 39)
NAVY2 = RGBColor(25, 52, 92)
NAVY3 = RGBColor(43, 77, 119)
GOLD = RGBColor(195, 161, 90)
PALE_GOLD = RGBColor(243, 227, 184)
IVORY = RGBColor(250, 248, 244)
BG = RGBColor(247, 248, 250)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(15, 23, 42)
MUTED = RGBColor(100, 116, 139)
GREEN = RGBColor(6, 118, 71)
BLUE = RGBColor(0, 122, 255)
LINE = RGBColor(226, 232, 240)
LIGHT_GREEN = RGBColor(236, 253, 243)
LIGHT_BLUE = RGBColor(230, 240, 252)
LIGHT_AMBER = RGBColor(255, 251, 235)

prs = Presentation(str(BASE))
blank = prs.slide_layouts[6]


def shape(slide, kind, x, y, w, h, fill=None, line=None, transparency=0):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
        shp.fill.transparency = transparency
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.8)
    return shp


def rounded(slide, x, y, w, h, fill=WHITE, line=None, transparency=0):
    return shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, transparency)


def rect(slide, x, y, w, h, fill, line=None, transparency=0):
    return shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, w, h, fill, line, transparency)


def circle(slide, x, y, d, fill, line=None):
    return shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, y, d, d, fill, line)


def text(slide, value, x, y, w, h, size=14, color=INK, bold=False, font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=11, color=INK, gap=5, bullet_color=GOLD):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.text = ""
        r = p.add_run()
        r.text = "•  "
        r.font.name = "Aptos"
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = bullet_color
        r2 = p.add_run()
        r2.text = item
        r2.font.name = "Aptos"
        r2.font.size = Pt(size)
        r2.font.color.rgb = color
    return box


def header(slide, kicker, title, subtitle, num, dark=False):
    title_color = WHITE if dark else INK
    muted = RGBColor(205, 216, 232) if dark else MUTED
    text(slide, kicker.upper(), 0.62, 0.43, 5.4, 0.22, 9.5, GOLD, True, margin=0)
    text(slide, title, 0.62, 0.70, 11.7, 0.58, 27, title_color, True, "Georgia", margin=0)
    text(slide, subtitle, 0.64, 1.32, 11.3, 0.34, 11.5, muted, margin=0)
    text(slide, f"{num:02d}", 12.35, 0.45, 0.4, 0.22, 9, muted, True, align=PP_ALIGN.RIGHT, margin=0)


def footer(slide, num, dark=False):
    color = RGBColor(185, 198, 218) if dark else RGBColor(148, 163, 184)
    rect(slide, 0.62, 7.13, 12.08, 0.008, GOLD if dark else LINE)
    text(slide, "Aurelia Collection · Hotel management project · 27 September 2026", 0.64, 7.19, 8.6, 0.18, 8.2, color, margin=0)
    text(slide, f"{num:02d}", 12.15, 7.16, 0.55, 0.22, 8.5, color, True, align=PP_ALIGN.RIGHT, margin=0)


def add_bg(slide, color=BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def card(slide, x, y, w, h, title, desc, fill=WHITE, accent=GOLD, title_color=INK, desc_color=MUTED):
    rounded(slide, x, y, w, h, fill, LINE if fill != NAVY else None)
    rect(slide, x, y, 0.055, h, accent)
    text(slide, title, x + 0.22, y + 0.20, w - 0.35, 0.24, 13, title_color, True, "Georgia", margin=0)
    text(slide, desc, x + 0.22, y + 0.62, w - 0.35, h - 0.70, 10.2, desc_color, margin=0)


# Slide 17 — project purpose and scope
slide = prs.slides.add_slide(blank)
add_bg(slide, IVORY)
header(slide, "Project purpose", "Why the hotel management project exists", "One system connects guest demand, hotel operations and financial control across a multi-property collection.", 17)
rounded(slide, 0.78, 1.95, 5.35, 4.70, NAVY, None)
text(slide, "The problem solved", 1.12, 2.30, 2.8, 0.24, 15, PALE_GOLD, True, "Georgia", margin=0)
text(slide, "A hotel should not need separate answers for booking, rooms, housekeeping and money.", 1.12, 2.77, 4.55, 0.75, 22, WHITE, True, "Georgia", margin=0)
bullets(slide, [
    "One booking record from any channel",
    "One availability and pricing engine",
    "One room-status and operations view",
    "One folio and financial history",
], 1.12, 4.10, 4.35, 1.45, 12, RGBColor(217, 227, 240), 9, PALE_GOLD)
rounded(slide, 6.48, 1.95, 6.14, 4.70, WHITE, LINE)
text(slide, "Project scope", 6.84, 2.30, 2.5, 0.24, 15, NAVY, True, "Georgia", margin=0)
items = [
    ("Guest", "Public collection, rooms, booking, payment intent, My Stays"),
    ("Front office", "Reservations, calendar, check-in, checkout, room assignment"),
    ("Operations", "Room Board, housekeeping, maintenance and live status"),
    ("Business", "Rates, folios, payments, reports, analytics and roles"),
    ("Deployment", "SQLite default, MySQL/MariaDB option, Navicat export"),
]
for i, (title, desc) in enumerate(items):
    y = 2.82 + i * 0.66
    circle(slide, 6.84, y + 0.03, 0.27, GOLD if i < 3 else NAVY3)
    text(slide, str(i + 1), 6.84, y + 0.10, 0.27, 0.10, 8, NAVY if i < 3 else WHITE, True, align=PP_ALIGN.CENTER, margin=0)
    text(slide, title, 7.28, y, 1.3, 0.18, 11, INK, True, margin=0)
    text(slide, desc, 8.62, y, 3.55, 0.24, 10, MUTED, margin=0)
footer(slide, 17)

# Slide 18 — data and system model
slide = prs.slides.add_slide(blank)
add_bg(slide, BG)
header(slide, "System model", "The hotel data that powers the PMS", "Every operating screen is built around properties, rooms, bookings, folios and people.", 18)
entities = [
    ("10", "Properties", "Hotels and destinations in the collection", NAVY),
    ("24", "Room types", "Sellable categories with base rates", GOLD),
    ("369", "Rooms", "Physical inventory assigned to room types", GREEN),
    ("170", "Bookings", "Reservations across all channels", NAVY3),
    ("170", "Folios", "One financial record for each stay", BLUE),
    ("103", "Payments", "Charges and refunds with status", RGBColor(130, 74, 20)),
]
for i, (num, title, desc, accent) in enumerate(entities):
    x = 0.78 + (i % 3) * 4.13
    y = 1.95 + (i // 3) * 1.55
    rounded(slide, x, y, 3.63, 1.17, WHITE, LINE)
    text(slide, num, x + 0.24, y + 0.22, 0.85, 0.38, 23, accent, True, "Georgia", margin=0)
    text(slide, title, x + 1.24, y + 0.23, 2.05, 0.20, 12.5, INK, True, "Georgia", margin=0)
    text(slide, desc, x + 1.24, y + 0.57, 2.08, 0.30, 9.6, MUTED, margin=0)
rounded(slide, 0.78, 5.36, 11.84, 1.05, NAVY, None)
text(slide, "Core relationship", 1.12, 5.66, 1.55, 0.18, 10, PALE_GOLD, True, margin=0)
text(slide, "Property  →  Room type  →  Room inventory  →  Booking  →  Folio  →  Payment", 2.85, 5.62, 8.95, 0.26, 17, WHITE, True, "Georgia", align=PP_ALIGN.CENTER, margin=0)
text(slide, "Users and roles control who can operate each part of this chain.", 3.16, 6.05, 8.30, 0.16, 9.5, RGBColor(200, 215, 234), align=PP_ALIGN.CENTER, margin=0)
footer(slide, 18)

# Slide 19 — website booking walkthrough
slide = prs.slides.add_slide(blank)
add_bg(slide, IVORY)
header(slide, "Guest flow", "Website booking in six simple steps", "The guest sees a clear path; the system performs the complex validation behind it.", 19)
steps = [
    ("1", "Discover", "View properties, maps, photos and hotel atmosphere."),
    ("2", "Search", "Choose dates, guests and rooms; availability is checked."),
    ("3", "Select", "Compare room types, amenities and nightly rates."),
    ("4", "Offer", "Eligible registered guests unlock a one-time 10% discount."),
    ("5", "Book", "Enter contact details, requests and payment method."),
    ("6", "Manage", "Receive confirmation, view My Stays and cancel if allowed."),
]
for i, (num, title, desc) in enumerate(steps):
    x = 0.78 + (i % 3) * 4.13
    y = 1.94 + (i // 3) * 2.06
    rounded(slide, x, y, 3.63, 1.55, WHITE, LINE)
    circle(slide, x + 0.25, y + 0.23, 0.42, NAVY if i in (0, 5) else GOLD)
    text(slide, num, x + 0.25, y + 0.35, 0.42, 0.12, 9, WHITE if i in (0, 5) else NAVY, True, align=PP_ALIGN.CENTER, margin=0)
    text(slide, title, x + 0.86, y + 0.26, 2.35, 0.22, 14, NAVY, True, "Georgia", margin=0)
    text(slide, desc, x + 0.25, y + 0.82, 3.05, 0.40, 10.2, MUTED, margin=0)
rounded(slide, 0.78, 6.28, 11.84, 0.43, LIGHT_GREEN, RGBColor(167, 243, 208))
text(slide, "Result: the guest gets a confirmation while the PMS receives a validated reservation, folio and financial snapshot.", 1.04, 6.41, 11.32, 0.16, 10, GREEN, True, align=PP_ALIGN.CENTER, margin=0)
footer(slide, 19)

# Slide 20 — daily operating playbook
slide = prs.slides.add_slide(blank)
add_bg(slide, BG)
header(slide, "Daily operations", "A simple playbook for the hotel team", "The PMS follows the property’s day: arrivals, rooms, guest service, departures and close.", 20)
playbook = [
    ("Morning", "Prepare", "Review arrivals, departures, occupancy, dirty rooms and open tasks.", GOLD),
    ("Arrival", "Check in", "Verify booking, assign vacant-clean room and move guest to in-house.", NAVY3),
    ("During stay", "Serve", "Monitor room status, requests, maintenance and folio activity.", GREEN),
    ("Departure", "Check out", "Settle folio, release room and automatically create housekeeping work.", BLUE),
    ("Night close", "Review", "Run nightly operations, flag no-shows and produce occupancy snapshot.", NAVY),
]
for i, (time, title, desc, accent) in enumerate(playbook):
    y = 1.92 + i * 0.90
    circle(slide, 0.92, y + 0.16, 0.40, accent)
    text(slide, str(i + 1), 0.92, y + 0.27, 0.40, 0.12, 9, WHITE, True, align=PP_ALIGN.CENTER, margin=0)
    rounded(slide, 1.60, y, 10.98, 0.68, WHITE, LINE)
    text(slide, time, 1.90, y + 0.14, 1.25, 0.18, 11, accent, True, "Georgia", margin=0)
    text(slide, title, 3.28, y + 0.14, 1.35, 0.18, 11, INK, True, margin=0)
    text(slide, desc, 4.84, y + 0.12, 7.25, 0.30, 10, MUTED, margin=0)
rounded(slide, 0.78, 6.52, 11.84, 0.33, NAVY, None)
text(slide, "Every action updates the same reservation, room status or folio — no disconnected spreadsheets required.", 1.05, 6.61, 11.28, 0.14, 9.8, WHITE, True, align=PP_ALIGN.CENTER, margin=0)
footer(slide, 20)

# Slide 21 — financial controls
slide = prs.slides.add_slide(blank)
add_bg(slide, NAVY)
header(slide, "Controls", "How the project protects the numbers", "The price shown, the folio issued and the historical total are designed to agree.", 21, dark=True)
card(slide, 0.78, 1.92, 3.70, 2.02, "Pricing engine", "Base rate + active rate rule + weekend rule, multiplied by nights and rooms.", RGBColor(23, 40, 68), GOLD, PALE_GOLD, RGBColor(210, 222, 238))
card(slide, 4.82, 1.92, 3.70, 2.02, "Charges", "The Cambodia demo uses 5% tax only; there is no separate service charge.", RGBColor(23, 40, 68), BLUE, PALE_GOLD, RGBColor(210, 222, 238))
card(slide, 8.86, 1.92, 3.70, 2.02, "Snapshot", "Rates, discount, service, tax and total are saved on the booking and never recomputed for history.", RGBColor(23, 40, 68), GREEN, PALE_GOLD, RGBColor(210, 222, 238))
rounded(slide, 0.78, 4.42, 7.62, 1.78, RGBColor(23, 40, 68), RGBColor(69, 92, 125))
text(slide, "Lifecycle safeguards", 1.12, 4.75, 2.5, 0.22, 14, PALE_GOLD, True, "Georgia", margin=0)
bullets(slide, [
    "Availability is checked across every night and protected against double booking.",
    "Cancellation rules can retain the first night and create a refund ledger entry.",
    "High or critical maintenance requests block a room from being sold.",
], 1.12, 5.16, 6.70, 0.82, 10.5, WHITE, 6, PALE_GOLD)
rounded(slide, 8.72, 4.42, 3.84, 1.78, PALE_GOLD, None)
text(slide, "QA proof", 9.04, 4.75, 1.8, 0.22, 14, NAVY, True, "Georgia", margin=0)
text(slide, "23/23 tests\n2,160 ledger invariants\n0 pending migrations", 9.04, 5.17, 3.08, 0.62, 15, GREEN, True, margin=0)
footer(slide, 21, dark=True)

# Slide 22 — delivery and decision summary
slide = prs.slides.add_slide(blank)
add_bg(slide, IVORY)
header(slide, "Final summary", "What the project delivers", "A clear hotel operating model ready for owner-controlled database configuration and future channel integration.", 22)
rounded(slide, 0.78, 1.95, 5.72, 4.65, NAVY, None)
text(slide, "The project is ready to demonstrate", 1.12, 2.32, 4.70, 0.25, 15, PALE_GOLD, True, "Georgia", margin=0)
text(slide, "From a guest’s first click to a manager’s revenue report, the important hotel actions are connected.", 1.12, 2.83, 4.55, 0.76, 21, WHITE, True, "Georgia", margin=0)
bullets(slide, [
    "Guest experience is public, responsive and premium.",
    "Staff workflow is organized by role and operating moment.",
    "Money is preserved as a historical contract.",
    "SQLite runs automatically; MySQL/Navicat is the deployment path.",
], 1.12, 4.25, 4.72, 1.45, 11.5, RGBColor(214, 225, 239), 8, PALE_GOLD)
rounded(slide, 6.82, 1.95, 5.80, 4.65, WHITE, LINE)
text(slide, "Final recommended sequence", 7.18, 2.32, 3.55, 0.25, 15, NAVY, True, "Georgia", margin=0)
final_steps = [
    ("1", "Use the public site", "Start signed out and test the guest journey."),
    ("2", "Test each role", "Confirm reception, housekeeping, manager and admin tasks."),
    ("3", "Configure MySQL", "Use the Navicat-owned server and private local settings."),
    ("4", "Add OTA integration", "Connect channel-manager APIs when ready."),
]
for i, (num, title, desc) in enumerate(final_steps):
    y = 2.90 + i * 0.68
    circle(slide, 7.18, y + 0.02, 0.30, GOLD)
    text(slide, num, 7.18, y + 0.10, 0.30, 0.10, 8, NAVY, True, align=PP_ALIGN.CENTER, margin=0)
    text(slide, title, 7.68, y, 1.75, 0.18, 10.8, INK, True, margin=0)
    text(slide, desc, 9.48, y, 2.55, 0.25, 9.7, MUTED, margin=0)
rounded(slide, 7.18, 5.82, 4.94, 0.48, LIGHT_GREEN, RGBColor(167, 243, 208))
text(slide, "Next action: verify the owner’s MySQL/Navicat connection.", 7.40, 5.98, 4.52, 0.15, 9.8, GREEN, True, align=PP_ALIGN.CENTER, margin=0)
footer(slide, 22)

prs.core_properties.title = "Aurelia Collection — Hotel Management Project Full Deck"
prs.core_properties.subject = "Complete hotel management project description, roles, channels, operations and controls"
prs.core_properties.author = "Aurelia Collection"
prs.core_properties.keywords = "hotel management, PMS, roles, website, walk-in, OTA, agency, operations, finance"
prs.core_properties.comments = "22-slide full project deck generated from the validated Aurelia Collection release."
prs.save(OUT)
print(OUT)
